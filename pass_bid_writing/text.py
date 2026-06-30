from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook


TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_legacy_doc_text(path: Path) -> str:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        raise ValueError(f"Reading .doc files requires Microsoft Word COM on Windows: {exc}") from exc

    word = None
    with tempfile.TemporaryDirectory() as tempdir:
        docx_path = Path(tempdir) / f"{path.stem}.docx"
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            document = word.Documents.Open(str(path.resolve()))
            document.SaveAs(str(docx_path), FileFormat=16)
            document.Close(False)
            return extract_docx_text(docx_path)
        except Exception as exc:
            raise ValueError(f"Failed to convert .doc file for text extraction: {exc}") from exc
        finally:
            if word is not None:
                try:
                    word.Quit()
                except Exception:
                    pass


def extract_pdf_text(path: Path, max_pages: int = 120) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages[:max_pages]:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts)


def extract_pdf_pages(path: Path, max_pages: int = 500) -> list[dict[str, Any]]:
    reader = PdfReader(str(path))
    pages: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages[:max_pages], start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append({"page": index, "text": text.strip()})
    return pages


def extract_located_text(path: Path) -> list[dict[str, Any]]:
    """Return source-located text blocks for evidence and reference extraction."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_pages(path)
    if suffix in TEXT_SUFFIXES or suffix in {".doc", ".docx"}:
        payload = extract_text(str(path))
        return [{"page": None, "text": payload["text"]}]
    if suffix == ".xlsx":
        return extract_xlsx_blocks(path)
    return []


def extract_xlsx_blocks(path: Path, max_rows_per_sheet: int = 5000) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(
                sheet.iter_rows(values_only=False),
                start=1,
            ):
                if row_index > max_rows_per_sheet:
                    break
                values = [
                    str(cell.value).strip()
                    for cell in row
                    if cell.value is not None and str(cell.value).strip()
                ]
                if not values:
                    continue
                blocks.append(
                    {
                        "page": None,
                        "sheet": sheet.title,
                        "cell": f"A{row_index}",
                        "text": " | ".join(values),
                    }
                )
    finally:
        workbook.close()
    return blocks


def extract_text(path_or_text: str) -> dict[str, Any]:
    candidate = Path(path_or_text).expanduser()
    if not candidate.exists():
        return {"source": "inline_text", "path": "", "text": path_or_text}

    path = candidate.resolve()
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        text = read_text_file(path)
    elif suffix == ".doc":
        text = extract_legacy_doc_text(path)
    elif suffix == ".docx":
        text = extract_docx_text(path)
    elif suffix == ".pdf":
        text = extract_pdf_text(path)
    elif suffix == ".xlsx":
        text = "\n".join(block["text"] for block in extract_xlsx_blocks(path))
    else:
        raise ValueError(f"Unsupported file type for text extraction: {path.suffix}")
    return {"source": "file", "path": str(path), "text": text}


def compact_text(text: str, limit: int = 600) -> str:
    clean = re.sub(r"\s+", " ", text).strip()
    if len(clean) <= limit:
        return clean
    return clean[: limit - 1].rstrip() + "…"


def split_meaningful_lines(text: str) -> list[tuple[int, str]]:
    lines: list[tuple[int, str]] = []
    for idx, raw in enumerate(text.splitlines(), start=1):
        line = re.sub(r"\s+", " ", raw).strip()
        if len(line) >= 4:
            lines.append((idx, line))
    return lines
