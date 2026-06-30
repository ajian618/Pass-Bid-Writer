from __future__ import annotations

from pathlib import Path
from typing import Any
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="E8F1FB")
HEADER_FONT = Font(name="微软雅黑", bold=True, color="17324D")
BODY_FONT = Font(name="微软雅黑", size=10)


def _write_workbook(path: Path, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column, value=_cell_value(value))
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row[column - 1] or "") for row in rows]
        width = min(max(max(len(value) for value in values) + 2, 10), 48)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 24
    workbook.save(path)


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return value


def export_production_reports(state: dict[str, Any], output_dir: Path) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, str] = {}

    task_path = output_dir / "施组编制任务书.xlsx"
    task_rows = []
    for section in state.get("sections", []):
        task_rows.append(
            [
                section.get("code"),
                section.get("title"),
                section.get("purpose"),
                "\n".join(section.get("required_inputs", [])),
                "\n".join(section.get("components", [])),
                "\n".join(section.get("acceptance", [])),
                "\n".join(section.get("project_basis", [])),
                "\n".join(section.get("reference_basis", [])),
                "\n".join(section.get("missing_inputs", [])),
                section.get("status"),
            ]
        )
    _write_workbook(
        task_path,
        "施组编制任务书",
        ["章节", "名称", "目的", "所需输入", "应交付组件", "完成标准", "项目依据", "编制依据", "缺失输入", "状态"],
        task_rows,
    )
    paths["task_spec"] = str(task_path)

    matrix_path = output_dir / "招标要求响应矩阵.xlsx"
    matrix_rows = [
        [
            row.get("requirement_id"),
            row.get("category"),
            row.get("requirement"),
            row.get("source_page") or row.get("source_line"),
            row.get("target_section"),
            row.get("response_strategy"),
            row.get("draft_status"),
            "是" if row.get("human_check") else "否",
        ]
        for row in state.get("response_matrix", [])
    ]
    _write_workbook(
        matrix_path,
        "招标要求响应矩阵",
        ["要求编号", "类别", "招标要求", "来源位置", "目标章节", "响应策略", "状态", "人工复核"],
        matrix_rows,
    )
    paths["response_matrix"] = str(matrix_path)

    basis_path = output_dir / "资料与规范依据报告.xlsx"
    basis_rows = []
    used_paths = {
        str(item.get("source_path", ""))
        for item in [*state.get("facts", []), *state.get("standards", [])]
        if item.get("source_path")
    }
    for file_info in state.get("files", []):
        basis_rows.append(
            [
                "项目资料",
                file_info.get("name"),
                file_info.get("role"),
                file_info.get("path"),
                "",
                "已利用" if file_info.get("path") in used_paths else "待利用/仅作为附件",
                file_info.get("sha256", ""),
            ]
        )
    for fact in state.get("facts", []):
        basis_rows.append(
            [
                "项目事实",
                fact.get("key"),
                fact.get("value"),
                fact.get("source_path"),
                fact.get("source_page"),
                fact.get("status"),
                fact.get("confidence"),
            ]
        )
    for standard in state.get("standards", []):
        basis_rows.append(
            [
                "规范标准",
                standard.get("code"),
                standard.get("title"),
                standard.get("source_path"),
                standard.get("source_page"),
                standard.get("status"),
                standard.get("official_url"),
            ]
        )
    for clause in state.get("standard_clauses", []):
        basis_rows.append(
            [
                "规范条款",
                f"{clause.get('standard_code')} {clause.get('clause_id')}",
                clause.get("text"),
                clause.get("source_path"),
                clause.get("page"),
                clause.get("verification_status"),
                clause.get("mandatory_level"),
            ]
        )
    for asset in state.get("case_assets", []):
        basis_rows.append(
            [
                "案例资产",
                asset.get("title"),
                asset.get("reuse_rule"),
                asset.get("source_path"),
                "",
                "可用",
                "\n".join(asset.get("applicable_sections", [])),
            ]
        )
    _write_workbook(
        basis_path,
        "资料与规范依据",
        ["依据类型", "名称或编号", "内容", "来源文件", "页码", "状态", "置信度或官方来源"],
        basis_rows,
    )
    paths["basis_report"] = str(basis_path)

    standards_path = output_dir / "规范待下载清单.xlsx"
    standard_rows = [
        [
            item.get("code"),
            item.get("title"),
            item.get("source_page"),
            item.get("official_platform"),
            item.get("search_query"),
            "\n".join(item.get("affected_sections", [])),
            item.get("status"),
        ]
        for item in state.get("standards", [])
        if item.get("status") in {"pending_download", "metadata_only", "blocked"}
    ]
    _write_workbook(
        standards_path,
        "规范待下载清单",
        ["标准编号", "标准名称", "招标引用页", "建议官方平台", "搜索关键词", "影响章节", "状态"],
        standard_rows,
    )
    paths["standards_queue"] = str(standards_path)

    confirmation_path = output_dir / "待补漏和冲突清单.xlsx"
    confirmation_rows = [
        [
            item.get("category"),
            item.get("title"),
            item.get("detail"),
            item.get("severity"),
            "\n".join(item.get("affected_sections", [])),
            item.get("status"),
            item.get("recommended_action", ""),
            item.get("resolution", ""),
        ]
        for item in state.get("confirmations", [])
    ]
    _write_workbook(
        confirmation_path,
        "待补漏和冲突",
        ["类别", "事项", "说明", "严重性", "影响章节", "状态", "建议动作", "处理结果"],
        confirmation_rows,
    )
    paths["confirmations"] = str(confirmation_path)

    differences_path = output_dir / "模型复核差异报告.xlsx"
    difference_rows = [
        [
            item.get("task"),
            item.get("primary_model"),
            item.get("review_model"),
            item.get("primary_result"),
            item.get("review_result"),
            item.get("difference"),
            item.get("status"),
        ]
        for item in state.get("model_differences", [])
    ]
    _write_workbook(
        differences_path,
        "模型复核差异",
        ["任务", "主模型", "复核模型", "主结果", "复核结果", "差异", "处理状态"],
        difference_rows,
    )
    paths["model_differences"] = str(differences_path)
    return paths
