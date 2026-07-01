from __future__ import annotations

import hashlib
import json
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from . import db
from .analysis import build_response_matrix, extract_tender_requirements
from .assets import generate_document_assets
from .blueprints import build_blueprint
from .config import ensure_storage_dirs, get_settings
from .documents import (
    check_docx_compliance,
    export_docx_to_pdf,
    generate_docx,
    safe_filename,
)
from .drawings import build_drawing_registry, render_confirmed_drawing
from .evidence_links import (
    is_reference_only_value,
    normalize_fact_value,
    parse_document_reference,
    strip_reference_tail,
)
from .knowledge import canonical_section_title, load_case_assets
from .model_router import ModelRouter
from .projects import expand_project_archives, scan_single_project
from .reports import export_production_reports
from .schemas import validate_core_state
from .text import compact_text, extract_located_text
from .visual_sources import build_visual_jobs
from .vision import visual_check_document


STANDARD_PATTERN = re.compile(
    r"(?P<code>(?:GB|GB/T|SL|SL/T|DL/T|JGJ|CJJ|JTG|NB/T|DB\d{2}/T?)\s*"
    r"\d+(?:\.\d+)?(?:[-—–]\d{4})?)",
    re.IGNORECASE,
)
FACT_PATTERNS = [
    ("总工期", re.compile(r"(?:计划)?总工期[^\d]{0,12}(\d+(?:\.\d+)?)\s*(日历天|天|个月)")),
    ("计划工期", re.compile(r"计划工期[^\d]{0,12}(\d+(?:\.\d+)?)\s*(日历天|天|个月)")),
    ("质量目标", re.compile(r"质量(?:目标|标准)[：:\s]*([^。；\n|]{2,40})")),
    ("建设地点", re.compile(r"(?:建设|工程|项目)地点[：:\s]*([^。；\n|]{2,60})")),
    ("项目名称", re.compile(r"(?:项目|工程)名称[：:\s]*([^。；\n|]{2,80})")),
]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def _official_reference(code: str) -> tuple[str, str]:
    upper = code.upper().replace(" ", "")
    if upper.startswith("SL"):
        return "水利部及水利行业标准平台", "https://www.mwr.gov.cn/"
    if upper.startswith("DB33"):
        return "浙江省政务公开及地方标准平台", "https://www.zj.gov.cn/"
    return "国家标准全文公开系统", "https://openstd.samr.gov.cn/"


def extract_standards(located_blocks: list[dict[str, Any]], source_path: str) -> list[dict[str, Any]]:
    standards: dict[str, dict[str, Any]] = {}
    for block in located_blocks:
        page = block.get("page")
        for raw_line in str(block.get("text", "")).splitlines():
            line = re.sub(r"\s+", " ", raw_line).strip()
            if not any(word in line for word in ("规范", "规程", "标准", "导则", "技术要求")):
                continue
            for match in STANDARD_PATTERN.finditer(line):
                code = re.sub(r"\s+", " ", match.group("code").upper()).strip()
                title_match = re.search(r"[《〈]([^》〉]{2,80})[》〉]", line)
                title = title_match.group(1).strip() if title_match else compact_text(line, 100)
                platform, official_url = _official_reference(code)
                standards.setdefault(
                    code,
                    {
                        "code": code,
                        "title": title,
                        "version": code.rsplit("-", 1)[-1] if re.search(r"[-—–]\d{4}$", code) else "",
                        "source_page": page,
                        "source_path": source_path,
                        "official_platform": platform,
                        "official_url": official_url,
                        "search_query": f"{code} {title} 官方 PDF",
                        "search_url": official_url,
                        "local_path": "",
                        "sha256": "",
                        "status": "pending_download",
                        "affected_sections": [],
                    },
                )
    return list(standards.values())


def extract_project_facts(files: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return extract_project_evidence(files)["facts"]


def extract_project_evidence(files: list[dict[str, Any]]) -> dict[str, Any]:
    facts: list[dict[str, Any]] = []
    reference_links: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, Any]] = set()
    for file_info in files:
        if file_info.get("is_duplicate"):
            continue
        path = Path(file_info["path"])
        if path.suffix.lower() not in {".pdf", ".doc", ".docx", ".txt", ".md", ".xlsx"}:
            continue
        try:
            blocks = extract_located_text(path)
        except Exception:
            continue
        for block in blocks:
            text = str(block.get("text", ""))
            for key, pattern in FACT_PATTERNS:
                for match in pattern.finditer(text):
                    raw_value = match.group(1).strip()
                    unit = match.group(2) if match.lastindex and match.lastindex >= 2 else ""
                    reference = parse_document_reference(match.group(0))
                    if reference and is_reference_only_value(raw_value):
                        reference_links.append(
                            {
                                "kind": "fact_reference",
                                "fact_key": key,
                                "source_path": str(path),
                                "source_page": block.get("page"),
                                "source_sheet": block.get("sheet", ""),
                                "source_cell": block.get("cell", ""),
                                "source_excerpt": compact_text(match.group(0), 260),
                                **reference,
                                "status": "unresolved",
                            }
                        )
                        continue
                    value = strip_reference_tail(raw_value)
                    if not value:
                        continue
                    identity = (key, value, str(path), block.get("page"))
                    if identity in seen:
                        continue
                    seen.add(identity)
                    facts.append(
                        {
                            "key": key,
                            "value": value,
                            "unit": unit,
                            "source_path": str(path),
                            "source_page": block.get("page"),
                            "source_sheet": block.get("sheet", ""),
                            "source_cell": block.get("cell", ""),
                            "source_excerpt": _fact_source_context(text, match.start(), match.end()),
                            "source_block_excerpt": compact_text(text, 800),
                            "confidence": 0.86,
                            "status": "extracted",
                        }
                    )
    _resolve_fact_references(facts, reference_links)
    merged_facts = _merge_equivalent_facts(facts)
    return {"facts": merged_facts, "reference_links": reference_links}


def _fact_source_context(text: str, start: int, end: int) -> str:
    return compact_text(text[max(0, start - 80) : min(len(text), end + 140)], 320)


def _resolve_fact_references(
    facts: list[dict[str, Any]],
    reference_links: list[dict[str, Any]],
) -> None:
    for reference in reference_links:
        candidates = [
            fact
            for fact in facts
            if fact.get("key") == reference.get("fact_key")
            and fact.get("source_path") == reference.get("source_path")
        ]
        target_page = reference.get("target_page")
        if target_page is not None:
            exact = [
                fact for fact in candidates if fact.get("source_page") == target_page
            ]
            if exact:
                candidates = exact
        target_table = str(reference.get("target_table", ""))
        if target_table:
            table_matches = [
                fact
                for fact in candidates
                if target_table
                in re.sub(r"\s+", "", str(fact.get("source_block_excerpt", "")))
            ]
            if table_matches:
                candidates = table_matches
        normalized_groups = {
            normalize_fact_value(
                str(fact.get("key", "")),
                str(fact.get("value", "")),
                str(fact.get("unit", "")),
            )
            for fact in candidates
        }
        if not candidates or len(normalized_groups) != 1:
            continue
        target = min(
            candidates,
            key=lambda fact: abs(
                int(fact.get("source_page") or 0) - int(target_page or 0)
            ),
        )
        reference["status"] = "resolved"
        reference["resolved_value"] = target.get("value", "")
        reference["resolved_page"] = target.get("source_page")
        target["status"] = "extracted"
        target["confidence"] = max(float(target.get("confidence", 0)), 0.95)
        target.setdefault("reference_chain", []).append(
            {
                "source_path": reference.get("source_path", ""),
                "source_page": reference.get("source_page"),
                "source_excerpt": reference.get("source_excerpt", ""),
                "target_page": target_page,
                "target_table": target_table,
            }
        )


def _merge_equivalent_facts(facts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    order: list[tuple[str, str]] = []
    for fact in facts:
        identity = (
            str(fact.get("key", "")),
            normalize_fact_value(
                str(fact.get("key", "")),
                str(fact.get("value", "")),
                str(fact.get("unit", "")),
            ),
        )
        if identity not in groups:
            groups[identity] = []
            order.append(identity)
        groups[identity].append(fact)
    merged: list[dict[str, Any]] = []
    for identity in order:
        items = groups[identity]
        primary = max(
            items,
            key=lambda item: (
                bool(item.get("reference_chain")),
                float(item.get("confidence", 0)),
                bool(item.get("source_page")),
            ),
        )
        primary = dict(primary)
        source_mentions = [
            {
                "source_path": item.get("source_path", ""),
                "source_page": item.get("source_page"),
                "source_sheet": item.get("source_sheet", ""),
                "source_cell": item.get("source_cell", ""),
                "source_excerpt": item.get("source_excerpt", ""),
            }
            for item in items
        ]
        primary["source_mentions"] = source_mentions
        primary["mention_count"] = len(source_mentions)
        primary["reference_chain"] = [
            link
            for item in items
            for link in item.get("reference_chain", [])
        ]
        merged.append(primary)
    return merged


def _detect_fact_conflicts(facts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    values: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for fact in facts:
        if not str(fact.get("value", "")).strip():
            continue
        normalized = normalize_fact_value(
            str(fact.get("key", "")),
            str(fact.get("value", "")),
            str(fact.get("unit", "")),
        )
        values.setdefault(fact["key"], {}).setdefault(normalized, []).append(fact)
    conflicts: list[dict[str, Any]] = []
    for key, groups in values.items():
        if len(groups) <= 1:
            continue
        detail = "；".join(
            f"{items[0].get('value', '')}{items[0].get('unit', '')}"
            f"（{Path(items[0]['source_path']).name}"
            f"{' 第' + str(items[0]['source_page']) + '页' if items[0].get('source_page') else ''}）"
            for items in groups.values()
        )
        conflicts.append(
            {
                "category": "资料冲突",
                "title": f"{key}存在不一致",
                "detail": detail,
                "severity": "high",
                "affected_sections": ["工程概况与施工条件", "施工进度计划及保证措施"],
                "status": "open",
                "recommended_action": "核对招标文件、补遗和设计资料并确认采用值。",
            }
        )
    return conflicts


def _map_standard_sections(standards: list[dict[str, Any]], sections: list[dict[str, Any]]) -> None:
    for standard in standards:
        blob = f"{standard.get('code', '')} {standard.get('title', '')}"
        affected: list[str] = []
        keyword_map = [
            (("质量", "验收", "检验"), ("质量", "材料", "试验")),
            (("安全", "施工现场", "临时用电"), ("安全", "文明", "应急")),
            (("水土保持", "环境"), ("环境", "水土保持")),
            (("土方", "堤防", "水利水电"), ("施工方案", "施工方法", "技术")),
        ]
        for keywords, section_keywords in keyword_map:
            if any(keyword in blob for keyword in keywords):
                matching = [
                    section["title"]
                    for section in sections
                    if any(keyword in section["title"] for keyword in section_keywords)
                ]
                affected.extend(matching[:2])
        if not affected:
            affected = [section["title"] for section in sections[:1]]
        standard["affected_sections"] = list(dict.fromkeys(affected))


def _build_sections(
    blueprint: dict[str, Any],
    requirements: dict[str, Any],
    matrix: list[dict[str, Any]],
    facts: list[dict[str, Any]],
    standards: list[dict[str, Any]],
    case_assets: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    for section in blueprint["sections"]:
        title = section["title"]
        matching_requirements = [
            row["requirement"]
            for row in matrix
            if canonical_section_title(str(row.get("target_section", ""))) == title
            or any(token in canonical_section_title(str(row.get("target_section", ""))) for token in title.split("、"))
        ]
        relevant_facts = [
            f"{fact['key']}：{fact['value']}{fact.get('unit', '')}"
            for fact in facts
            if any(keyword in title for keyword in _fact_section_keywords(fact["key"]))
        ]
        relevant_standards = [
            f"{standard['code']} {standard['title']}"
            for standard in standards
            if title in standard.get("affected_sections", [])
        ]
        project_basis = matching_requirements[:8] + relevant_facts[:8]
        relevant_case_assets = [
            asset
            for asset in case_assets
            if "全局" in asset.get("applicable_sections", [])
            or title in asset.get("applicable_sections", [])
        ]
        case_basis = [
            f"案例资产：{asset['case_title']} / {asset['title']}（{asset['reuse_rule']}）"
            for asset in relevant_case_assets[:3]
        ]
        reference_basis = (
            [f"蓝图：{blueprint['title']} {blueprint['version']}"]
            + relevant_standards[:8]
            + case_basis
        )
        missing_inputs = [
            item
            for item in section["required_inputs"]
            if not _input_is_available(item, facts, requirements, standards)
        ]
        completion = max(5, min(90, 20 + len(project_basis) * 8 + len(reference_basis) * 5 - len(missing_inputs) * 5))
        sections.append(
            {
                "code": section["code"],
                "title": title,
                "purpose": section["purpose"],
                "required_inputs": section["required_inputs"],
                "components": section["components"],
                "acceptance": section["acceptance"],
                "requirements": matching_requirements,
                "project_basis": project_basis,
                "reference_basis": reference_basis,
                "case_asset_ids": [asset["asset_id"] for asset in relevant_case_assets[:6]],
                "missing_inputs": missing_inputs,
                "completion": completion,
                "status": "ready" if not missing_inputs else "needs_input",
                "model": "deepseek-v4-pro",
            }
        )
    return sections


def _fact_section_keywords(fact_key: str) -> tuple[str, ...]:
    return {
        "总工期": ("进度", "部署", "季节"),
        "计划工期": ("进度", "部署", "季节"),
        "质量目标": ("质量", "编制"),
        "建设地点": ("概况", "总平面"),
        "项目名称": ("编制", "概况"),
    }.get(fact_key, ("概况",))


def _input_is_available(
    input_name: str,
    facts: list[dict[str, Any]],
    requirements: dict[str, Any],
    standards: list[dict[str, Any]],
) -> bool:
    blob = " ".join(
        [input_name]
        + [str(item.get("requirement", "")) for item in requirements.get("requirements", [])]
        + [f"{fact.get('key', '')}{fact.get('value', '')}" for fact in facts]
        + [f"{item.get('code', '')}{item.get('title', '')}" for item in standards]
    )
    keyword_groups = {
        "招标文件": ("招标", "技术标", "施工组织设计"),
        "引用规范": ("规范", "标准", "规程"),
        "总工期": ("工期",),
        "质量目标": ("质量",),
    }
    keywords = keyword_groups.get(input_name, tuple(re.findall(r"[\u4e00-\u9fff]{2,}", input_name)))
    return any(keyword in blob and any(
        keyword in str(value)
        for value in (
            [item.get("requirement", "") for item in requirements.get("requirements", [])]
            + [f"{fact.get('key', '')}{fact.get('value', '')}" for fact in facts]
            + [f"{item.get('code', '')}{item.get('title', '')}" for item in standards]
        )
    ) for keyword in keywords)


def _build_confirmations(
    sections: list[dict[str, Any]],
    standards: list[dict[str, Any]],
    conflicts: list[dict[str, Any]],
    files: list[dict[str, Any]],
    case_assets: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    confirmations = list(conflicts)
    for standard in standards:
        if standard["status"] == "pending_download":
            confirmations.append(
                {
                    "category": "规范待下载",
                    "title": f"缺少规范原文：{standard['code']}",
                    "detail": standard["title"],
                    "severity": "medium",
                    "affected_sections": standard["affected_sections"],
                    "status": "open",
                    "recommended_action": f"从{standard['official_platform']}下载并上传原文。",
                }
            )
    missing_roles = {
        "design_report": "初步设计或批复报告",
        "budget": "预算或工程量清单",
        "drawing": "设计图纸或CAD图纸包",
    }
    available_roles = {item.get("role") for item in files}
    for role, label in missing_roles.items():
        if role not in available_roles:
            confirmations.append(
                {
                    "category": "资料缺失",
                    "title": f"未识别到{label}",
                    "detail": "系统无法完成相应项目事实、工程量或图纸引用核验。",
                    "severity": "high" if role != "drawing" else "medium",
                    "affected_sections": [
                        section["title"]
                        for section in sections
                        if any(label_part in " ".join(section["required_inputs"]) for label_part in label.split("或"))
                    ][:6],
                    "status": "open",
                    "recommended_action": f"补充上传{label}并重新分析。",
                }
            )
    if not case_assets:
        confirmations.append(
            {
                "category": "案例资产缺失",
                "title": "当前没有可用的公司已通过成品",
                "detail": "本次任务书采用受控水利通用蓝图，不引用历史案例措辞或参数。",
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "导入1—2套招标文件与已通过技术标案例后重新建立任务。",
            }
        )
    for section in sections:
        if section["missing_inputs"]:
            confirmations.append(
                {
                    "category": "章节输入缺失",
                    "title": f"{section['title']}尚缺输入",
                    "detail": "、".join(section["missing_inputs"]),
                    "severity": "medium",
                    "affected_sections": [section["title"]],
                    "status": "open",
                    "recommended_action": "补充资料或人工确认后再生成正式正文。",
                }
            )
    return confirmations


def _build_reference_confirmations(
    reference_links: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    confirmations: list[dict[str, Any]] = []
    for item in reference_links:
        if item.get("status") == "resolved":
            continue
        target_parts = []
        if item.get("target_page"):
            target_parts.append(f"第{item['target_page']}页")
        if item.get("target_table"):
            target_parts.append(str(item["target_table"]))
        target = " ".join(target_parts) or "未明确位置"
        category = (
            "项目事实引用待解析"
            if item.get("kind") == "fact_reference"
            else "招标要求引用待解析"
        )
        confirmations.append(
            {
                "category": category,
                "title": f"未找到引用目标：{item.get('fact_key') or item.get('category', '招标要求')}",
                "detail": (
                    f"{item.get('source_excerpt', '')}；指向：{target}。"
                    "系统未把引用语当作实际项目值。"
                ),
                "severity": "high" if item.get("kind") == "fact_reference" else "medium",
                "affected_sections": ["编制说明及工程概况"],
                "status": "open",
                "recommended_action": "核对目标页或附表名称，补充可检索文件后重新分析。",
            }
        )
    return confirmations


def prepare_production_project(
    project_dir: str,
    *,
    project_type: str = "水利工程通用",
    expand_archives: bool = True,
    role_overrides: dict[str, str] | None = None,
) -> dict[str, Any]:
    settings = get_settings()
    ensure_storage_dirs(settings)
    db.init_db(settings.database_path)
    project_path = Path(project_dir).expanduser()
    if not project_path.is_absolute():
        project_path = settings.root_dir / project_path
    project_path = project_path.resolve()
    archive_result = expand_project_archives(project_path) if expand_archives else {"extracted": [], "skipped": []}
    project = scan_single_project(project_path, project_kind="new_tender")
    if role_overrides:
        normalized_overrides = {
            str(Path(path).expanduser().resolve()): role
            for path, role in role_overrides.items()
        }
        for file_info in project["files"]:
            override = normalized_overrides.get(str(Path(file_info["path"]).resolve()))
            if override:
                file_info["role"] = override
                file_info["confidence"] = 1.0
                file_info["reason"] = "human confirmed role"
                file_info["role_confirmed"] = True
    tender_files = [
        item
        for item in project["files"]
        if item["role"] in {"tender", "candidate_tender"} and not item.get("is_duplicate")
    ]
    if not tender_files:
        raise ValueError("项目中没有可识别的招标文件")
    tender_file = sorted(tender_files, key=lambda item: -float(item["confidence"]))[0]
    tender_blocks = extract_located_text(Path(tender_file["path"]))
    tender_text = "\n".join(block["text"] for block in tender_blocks)
    requirements = extract_tender_requirements(
        tender_text,
        source_path=tender_file["path"],
        located_blocks=tender_blocks,
    )
    matrix = build_response_matrix(requirements)
    project_evidence = extract_project_evidence(project["files"])
    facts = project_evidence["facts"]
    evidence_links = [
        *requirements.get("reference_links", []),
        *project_evidence.get("reference_links", []),
    ]
    standards = extract_standards(tender_blocks, tender_file["path"])
    blueprint = build_blueprint(project_type)
    _map_standard_sections(standards, blueprint["sections"])
    requirement_text = " ".join(
        str(item.get("requirement", "")) for item in requirements.get("requirements", [])
    )
    case_assets = load_case_assets(project_type=project_type, requirement_text=requirement_text)
    sections = _build_sections(blueprint, requirements, matrix, facts, standards, case_assets)
    conflicts = _detect_fact_conflicts(facts)
    confirmations = _build_confirmations(
        sections,
        standards,
        conflicts,
        project["files"],
        case_assets,
    )
    confirmations.extend(_build_reference_confirmations(evidence_links))
    model_router = ModelRouter()
    visual_jobs = build_visual_jobs(project["files"], project_root=project_path)
    drawings, drawing_confirmations = build_drawing_registry(project["files"])
    confirmations.extend(drawing_confirmations)
    state: dict[str, Any] = {
        "project": {
            "name": project["name"],
            "project_dir": project["project_dir"],
            "project_type": project_type,
            "file_count": project["file_count"],
        },
        "stage": "角色确认",
        "stage_index": 2,
        "status": "awaiting_role_confirmation",
        "file_roles_confirmed": False,
        "archive_result": archive_result,
        "files": project["files"],
        "requirements": requirements,
        "response_matrix": matrix,
        "facts": facts,
        "evidence_links": evidence_links,
        "standards": standards,
        "standard_clauses": [],
        "blueprint": blueprint,
        "case_assets": case_assets,
        "drawings": drawings,
        "visual_jobs": visual_jobs,
        "visual_analysis": {
            "status": "pending",
            "total": len(visual_jobs),
            "completed": 0,
            "failed": 0,
            "needs_review": 0,
        },
        "sections": sections,
        "task_spec_source": "company_compact_blueprint",
        "task_spec_rationale": (
            "先以公司常用的8章紧凑结构建立可编辑任务书；"
            "完成资料与案例匹配后可重新生成建议目录，最终以人工确认版本为准。"
        ),
        "task_spec_revision": 1,
        "task_spec_history": [],
        "confirmations": confirmations,
        "model_differences": [],
        "models": model_router.status(),
        "metrics": _metrics(
            project["files"],
            requirements,
            facts,
            standards,
            sections,
            confirmations,
            evidence_links,
        ),
        "created_at": _now(),
    }
    validate_core_state(state)
    with db.db_session(settings.database_path) as conn:
        project_id = db.create_or_update_project(
            conn,
            name=project["name"],
            project_kind="new_tender",
            project_dir=project["project_dir"],
            summary=project,
        )
        for file_info in project["files"]:
            db.upsert_project_file(
                conn,
                project_id=project_id,
                file_path=file_info["path"],
                file_name=file_info["name"],
                suffix=file_info["suffix"],
                role=file_info["role"],
                confidence=float(file_info["confidence"]),
                metadata={"sha256": file_info.get("sha256", ""), "reason": file_info.get("reason", "")},
            )
        now = _now()
        cursor = conn.execute(
            """
            INSERT INTO production_runs
            (project_id, project_type, stage, status, state_json, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (project_id, project_type, state["stage"], state["status"], _json(state), now, now),
        )
        run_id = int(cursor.lastrowid)
        state["run_id"] = run_id
        _persist_run_entities(conn, run_id, state)
        conn.execute("UPDATE production_runs SET state_json = ? WHERE id = ?", (_json(state), run_id))
    output_dir = project_path / "outputs"
    reports = export_production_reports(state, output_dir)
    state["reports"] = reports
    with db.db_session(settings.database_path) as conn:
        conn.execute(
            "UPDATE production_runs SET state_json = ?, updated_at = ? WHERE id = ?",
            (_json(state), _now(), state["run_id"]),
        )
    return state


def _attach_requirement_locations(
    requirements: dict[str, Any],
    located_blocks: list[dict[str, Any]],
) -> None:
    for item in requirements.get("requirements", []):
        text = str(item.get("requirement", "")).strip()
        needle = re.sub(r"\s+", "", text[:36])
        for block in located_blocks:
            haystack = re.sub(r"\s+", "", str(block.get("text", "")))
            if needle and needle in haystack:
                item["source_page"] = block.get("page")
                item["source_path"] = requirements.get("source_path", "")
                break


def _persist_run_entities(conn, run_id: int, state: dict[str, Any]) -> None:
    now = _now()
    for standard in state["standards"]:
        conn.execute(
            """
            INSERT INTO standard_documents
            (production_run_id, standard_code, title, version, source_page, source_path,
             official_url, local_path, sha256, status, affected_sections_json, metadata_json,
             created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id, standard["code"], standard["title"], standard["version"],
                standard["source_page"], standard["source_path"], standard["official_url"],
                standard["local_path"], standard["sha256"], standard["status"],
                _json(standard["affected_sections"]),
                _json({"official_platform": standard["official_platform"], "search_query": standard["search_query"]}),
                now, now,
            ),
        )
    for clause in state.get("standard_clauses", []):
        conn.execute(
            """
            INSERT INTO standard_clauses
            (production_run_id, standard_code, clause_id, clause_text, source_path, source_page,
             applicable_sections_json, applicability, mandatory_level, verification_status,
             created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                clause["standard_code"],
                clause["clause_id"],
                clause["text"],
                clause["source_path"],
                clause.get("page"),
                _json(clause.get("applicable_sections", [])),
                clause.get("applicability", ""),
                clause.get("mandatory_level", "unknown"),
                clause.get("verification_status", "verified_local_source"),
                now,
                now,
            ),
        )
    for asset in state.get("case_assets", []):
        conn.execute(
            """
            INSERT INTO case_assets
            (production_run_id, case_id, asset_key, asset_type, title, content_json,
             applicable_sections_json, reuse_rule, source_path, layout_profile_id, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                asset["case_id"],
                asset["asset_id"],
                asset["asset_type"],
                asset["title"],
                _json(asset.get("content")),
                _json(asset.get("applicable_sections", [])),
                asset.get("reuse_rule", "structure_only"),
                asset.get("source_path", ""),
                asset.get("layout_profile_id"),
                now,
            ),
        )
    for fact in state["facts"]:
        conn.execute(
            """
            INSERT INTO project_facts
            (production_run_id, fact_key, value, unit, source_path, source_page,
             source_excerpt, confidence, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id, fact["key"], fact["value"], fact["unit"], fact["source_path"],
                fact["source_page"], fact["source_excerpt"], fact["confidence"], fact["status"], now,
            ),
        )
    for drawing in state.get("drawings", []):
        conn.execute(
            """
            INSERT INTO drawing_assets
            (production_run_id, source_dwg_path, source_pdf_path, source_page, drawing_no,
             title, caption, crop_json, applicable_sections_json, placement, preview_path,
             confidence, status, confirmed_at, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                drawing.get("source_dwg_path", ""),
                drawing.get("source_pdf_path", ""),
                drawing.get("source_page"),
                drawing.get("drawing_no", ""),
                drawing.get("title", ""),
                drawing.get("caption", ""),
                _json(drawing.get("crop", {})),
                _json(drawing.get("applicable_sections", [])),
                drawing.get("placement", "inline"),
                drawing.get("preview_path", ""),
                float(drawing.get("confidence", 0)),
                drawing.get("status", "pending"),
                drawing.get("confirmed_at", ""),
                now,
                now,
            ),
        )
    for section in state["sections"]:
        _insert_section_entity(conn, run_id, section, now)
    for item in state["confirmations"]:
        conn.execute(
            """
            INSERT INTO confirmation_items
            (production_run_id, category, title, detail, severity, affected_sections_json,
             status, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id, item["category"], item["title"], item["detail"], item["severity"],
                _json(item["affected_sections"]), item["status"], now, now,
            ),
        )


def _replace_section_entities(
    conn,
    run_id: int,
    sections: list[dict[str, Any]],
) -> None:
    conn.execute("DELETE FROM content_components WHERE production_run_id = ?", (run_id,))
    conn.execute("DELETE FROM section_tasks WHERE production_run_id = ?", (run_id,))
    now = _now()
    for section in sections:
        _insert_section_entity(conn, run_id, section, now)


def _replace_case_asset_entities(
    conn,
    run_id: int,
    assets: list[dict[str, Any]],
) -> None:
    conn.execute("DELETE FROM case_assets WHERE production_run_id = ?", (run_id,))
    now = _now()
    for asset in assets:
        conn.execute(
            """
            INSERT INTO case_assets
            (production_run_id, case_id, asset_key, asset_type, title, content_json,
             applicable_sections_json, reuse_rule, source_path, layout_profile_id, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                asset["case_id"],
                asset["asset_id"],
                asset["asset_type"],
                asset["title"],
                _json(asset.get("content")),
                _json(asset.get("applicable_sections", [])),
                asset.get("reuse_rule", "structure_only"),
                asset.get("source_path", ""),
                asset.get("layout_profile_id"),
                now,
            ),
        )


def _insert_section_entity(
    conn,
    run_id: int,
    section: dict[str, Any],
    now: str,
) -> None:
    conn.execute(
        """
        INSERT INTO section_tasks
        (production_run_id, section_code, title, purpose, status, completion,
         requirements_json, inputs_json, acceptance_json, components_json, basis_json,
         content_text, model, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            run_id,
            section["code"],
            section["title"],
            section["purpose"],
            section["status"],
            section["completion"],
            _json(section["requirements"]),
            _json(section["required_inputs"]),
            _json(section["acceptance"]),
            _json(section["components"]),
            _json(
                {
                    "project": section["project_basis"],
                    "reference": section["reference_basis"],
                }
            ),
            section.get("content", ""),
            section.get("model", ""),
            now,
            now,
        ),
    )
    for index, component in enumerate(section.get("components", []), start=1):
        conn.execute(
            """
            INSERT INTO content_components
            (production_run_id, section_code, component_id, component_kind, title,
             required, status, source_path, basis_json, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                section["code"],
                f"{section['code']}-{index:02d}",
                _component_kind(component),
                component,
                1,
                "generated" if section.get("status") == "generated" else "not_started",
                "",
                _json(
                    {
                        "project": section.get("project_basis", []),
                        "reference": section.get("reference_basis", []),
                    }
                ),
                now,
                now,
            ),
        )


def _metrics(
    files: list[dict[str, Any]],
    requirements: dict[str, Any],
    facts: list[dict[str, Any]],
    standards: list[dict[str, Any]],
    sections: list[dict[str, Any]],
    confirmations: list[dict[str, Any]],
    evidence_links: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    requirement_count = len(requirements.get("requirements", []))
    requirement_texts = {
        str(item.get("requirement", "")).strip()
        for item in requirements.get("requirements", [])
        if str(item.get("requirement", "")).strip()
    }
    mapped_requirements = {
        str(requirement).strip()
        for section in sections
        for requirement in section.get("requirements", [])
        if str(requirement).strip()
    }
    covered = len(requirement_texts.intersection(mapped_requirements))
    downloaded = sum(
        1 for standard in standards if standard["status"] in {"available", "not_applicable"}
    )
    avg_completion = round(sum(section["completion"] for section in sections) / max(len(sections), 1))
    evidence_links = evidence_links or []
    resolved_links = sum(
        1 for item in evidence_links if item.get("status") == "resolved"
    )
    return {
        "requirement_count": requirement_count,
        "requirement_coverage": min(100, round(covered / max(requirement_count, 1) * 100)),
        "file_count": len(files),
        "fact_count": len(facts),
        "standards_count": len(standards),
        "standards_ready": downloaded,
        "standards_readiness": round(downloaded / max(len(standards), 1) * 100) if standards else 100,
        "section_count": len(sections),
        "chapter_completion": avg_completion,
        "open_confirmations": sum(1 for item in confirmations if item["status"] == "open"),
        "high_risks": sum(1 for item in confirmations if item["severity"] == "high"),
        "reference_link_count": len(evidence_links),
        "reference_link_resolved": resolved_links,
        "reference_link_readiness": (
            round(resolved_links / len(evidence_links) * 100)
            if evidence_links
            else 100
        ),
    }


def _component_kind(title: str) -> str:
    if "甘特" in title or "进度" in title and "图" in title:
        return "gantt"
    if "组织" in title and "图" in title:
        return "organization_chart"
    if "流程" in title or "图" in title:
        return "flowchart"
    if "表" in title or "清单" in title:
        return "table"
    if "CAD" in title or "总平" in title:
        return "cad_reference"
    if "附件" in title:
        return "attachment"
    return "text"


def get_run_state(run_id: int | None = None) -> dict[str, Any] | None:
    settings = get_settings()
    ensure_storage_dirs(settings)
    db.init_db(settings.database_path)
    with db.db_session(settings.database_path) as conn:
        if run_id is None:
            row = conn.execute("SELECT * FROM production_runs ORDER BY id DESC LIMIT 1").fetchone()
        else:
            row = conn.execute("SELECT * FROM production_runs WHERE id = ?", (run_id,)).fetchone()
    if not row:
        return None
    state = json.loads(row["state_json"])
    for fact in state.get("facts", []):
        if fact.get("status") == "resolved_reference":
            fact["status"] = "extracted"
    state["run_id"] = int(row["id"])
    return state


def confirm_task_spec(run_id: int) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if not state.get("file_roles_confirmed"):
        raise ValueError("必须先确认项目文件角色，避免以错误资料生成正文")
    if state.get("status") == "visual_analysis_pending":
        raise ValueError("必须先运行多模态资料分析；未配置视觉模型时系统会生成明确的人工核对项")
    sections = state.get("sections", [])
    if not sections:
        raise ValueError("编制任务书至少需要一个章节")
    invalid = [
        section
        for section in sections
        if not str(section.get("title", "")).strip()
        or not section.get("components")
        or not section.get("acceptance")
    ]
    if invalid:
        raise ValueError("每个章节都必须填写章节名称、交付组件和完成标准")
    blocking = [
        item
        for item in state.get("confirmations", [])
        if item.get("status") == "open"
        and item.get("category") in {"资料冲突", "视觉资料冲突"}
        and item.get("severity") == "high"
    ]
    open_differences = [
        item for item in state.get("model_differences", []) if item.get("status") == "open"
    ]
    if blocking or open_differences:
        raise ValueError("存在未解决的项目数据冲突或模型分歧，相关章节不得定稿")
    state["stage"] = "章节生成"
    state["stage_index"] = 7
    state["status"] = "ready_for_generation"
    state["task_spec_confirmed_at"] = _now()
    state["task_spec_confirmed_revision"] = int(state.get("task_spec_revision", 1))
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        conn.execute(
            """
            UPDATE production_runs SET stage = ?, status = ?, state_json = ?, updated_at = ?
            WHERE id = ?
            """,
            (state["stage"], state["status"], _json(state), _now(), run_id),
        )
    return state


def suggest_task_spec(run_id: int) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    _ensure_task_spec_editable(state)
    blueprint = state.get("blueprint") or build_blueprint(
        state.get("project", {}).get("project_type", "水利工程通用")
    )
    requirement_text = " ".join(
        str(item.get("requirement", ""))
        for item in state.get("requirements", {}).get("requirements", [])
    )
    state["case_assets"] = load_case_assets(
        project_type=state.get("project", {}).get("project_type", ""),
        requirement_text=requirement_text,
    )
    state["confirmations"] = [
        item
        for item in state.get("confirmations", [])
        if item.get("category") != "案例资产缺失"
    ]
    if not state["case_assets"]:
        state["confirmations"].append(
            {
                "category": "案例资产缺失",
                "title": "当前没有已启用的公司已通过成品",
                "detail": "建议目录将采用受控水利通用蓝图，停用案例不会参与匹配。",
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "在成品蓝图中审核并启用适用案例，或继续人工编辑任务书。",
            }
        )
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        _replace_case_asset_entities(conn, run_id, state.get("case_assets", []))
    case_outlines = _case_outline_candidates(state.get("case_assets", []))
    router = ModelRouter()
    proposed: list[dict[str, Any]] = []
    rationale = ""
    source = "company_compact_blueprint"
    if router.endpoints["text_master"].configured:
        result = router.complete_json(
            role="text_master",
            system=(
                "你是浙江水利工程通过制技术标的编制任务书规划器。"
                "只规划一级章节和每章交付定义，不撰写正文，不补造任何项目参数或规范条文。"
                "章节数必须由招标要求、项目资料和已通过案例结构共同决定，不得固定为15章。"
                "输出严格JSON对象。"
            ),
            prompt=_task_spec_prompt(state, blueprint, case_outlines),
            max_tokens=7000,
        )
        proposed = result.get("sections", []) if isinstance(result, dict) else []
        rationale = str(result.get("rationale", "")).strip()
        source = "deepseek_evidence_and_case_match"
    if not proposed:
        titles = next(
            (
                item["titles"]
                for item in case_outlines
                if 4 <= len(item.get("titles", [])) <= 15
            ),
            [],
        )
        if titles:
            proposed = [{"title": title} for title in titles]
            source = "matched_case_outline"
            rationale = "采用已启用案例的一级目录作为建议起点，并按本项目证据重新计算章节依据和缺口。"
        else:
            proposed = blueprint.get("sections", [])
            rationale = "没有可用的稳定案例一级目录，保留公司常用8章蓝图作为可编辑起点。"
    return _apply_task_spec(
        run_id,
        state,
        proposed,
        source=source,
        rationale=rationale,
        revision_note="根据当前项目资料、招标要求和已启用案例重新生成建议目录",
    )


def update_task_spec(
    run_id: int,
    sections: list[dict[str, Any]],
    *,
    revision_note: str = "",
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    _ensure_task_spec_editable(state)
    return _apply_task_spec(
        run_id,
        state,
        sections,
        source="human_edited",
        rationale=revision_note.strip() or "人工修订章节结构和每章交付定义。",
        revision_note=revision_note.strip() or "人工修订任务书",
    )


def _ensure_task_spec_editable(state: dict[str, Any]) -> None:
    if not state.get("file_roles_confirmed"):
        raise ValueError("必须先确认项目文件角色")
    if state.get("status") in {
        "generating",
        "partial_draft",
        "ready_for_assembly",
        "draft_generated",
        "reviewed",
    } or any(section.get("content") for section in state.get("sections", [])):
        raise ValueError("正文生产已经开始，不能再改任务书；请新建生产任务后调整目录")


def _apply_task_spec(
    run_id: int,
    state: dict[str, Any],
    proposed: list[dict[str, Any]],
    *,
    source: str,
    rationale: str,
    revision_note: str,
) -> dict[str, Any]:
    if not isinstance(proposed, list) or not 1 <= len(proposed) <= 30:
        raise ValueError("任务书章节数必须在1至30章之间")
    titles: set[str] = set()
    definitions: list[dict[str, Any]] = []
    blueprint = state.get("blueprint") or build_blueprint(
        state.get("project", {}).get("project_type", "水利工程通用")
    )
    for index, raw in enumerate(proposed, start=1):
        if not isinstance(raw, dict):
            continue
        title = _clean_section_title(str(raw.get("title", "")))
        if not title:
            continue
        if title in titles:
            raise ValueError(f"章节名称重复：{title}")
        titles.add(title)
        template = _closest_section_template(title, blueprint)
        definitions.append(
            {
                "code": f"{index:02d}",
                "title": title,
                "purpose": str(raw.get("purpose", "")).strip()
                or template.get("purpose", "明确本章响应范围和可检查交付成果。"),
                "required_inputs": _clean_string_list(
                    raw.get("required_inputs") or template.get("required_inputs", [])
                ),
                "components": _clean_string_list(
                    raw.get("components") or template.get("components", [])
                ),
                "acceptance": _clean_string_list(
                    raw.get("acceptance") or template.get("acceptance", [])
                ),
            }
        )
    if not definitions:
        raise ValueError("任务书没有有效章节")
    if any(not item["components"] or not item["acceptance"] for item in definitions):
        raise ValueError("每个章节至少需要一个交付组件和一项完成标准")
    outline = [
        {"level": 1, "title": item["title"], "order": index}
        for index, item in enumerate(definitions, start=1)
    ]
    matrix = build_response_matrix(state.get("requirements", {}), outline=outline)
    _map_standard_sections(state.get("standards", []), definitions)
    dynamic_blueprint = {
        **blueprint,
        "sections": definitions,
        "version": f"{blueprint.get('version', '1.0')}-task",
    }
    sections = _build_sections(
        dynamic_blueprint,
        state.get("requirements", {}),
        matrix,
        state.get("facts", []),
        state.get("standards", []),
        state.get("case_assets", []),
    )
    previous_snapshot = [
        {
            "code": item.get("code", ""),
            "title": item.get("title", ""),
            "purpose": item.get("purpose", ""),
            "required_inputs": item.get("required_inputs", []),
            "components": item.get("components", []),
            "acceptance": item.get("acceptance", []),
        }
        for item in state.get("sections", [])
    ]
    state.setdefault("task_spec_history", []).append(
        {
            "revision": int(state.get("task_spec_revision", 1)),
            "saved_at": _now(),
            "source": state.get("task_spec_source", ""),
            "note": revision_note,
            "sections": previous_snapshot,
        }
    )
    state["sections"] = sections
    state["response_matrix"] = matrix
    state["task_spec_source"] = source
    state["task_spec_rationale"] = rationale
    state["task_spec_revision"] = int(state.get("task_spec_revision", 1)) + 1
    state["task_spec_confirmed_at"] = ""
    state["task_spec_confirmed_revision"] = None
    state["stage"] = "编制任务书确认"
    state["stage_index"] = 6
    state["status"] = "awaiting_confirmation"
    state["confirmations"] = [
        item
        for item in state.get("confirmations", [])
        if item.get("category") != "章节输入缺失"
    ]
    for section in sections:
        if section.get("missing_inputs"):
            state["confirmations"].append(
                {
                    "category": "章节输入缺失",
                    "title": f"{section['title']}尚缺输入",
                    "detail": "、".join(section["missing_inputs"]),
                    "severity": "medium",
                    "affected_sections": [section["title"]],
                    "status": "open",
                    "recommended_action": "补充资料或人工确认后再生成正式正文。",
                }
            )
    state["metrics"] = _metrics(
        state.get("files", []),
        state.get("requirements", {}),
        state.get("facts", []),
        state.get("standards", []),
        sections,
        state.get("confirmations", []),
        state.get("evidence_links", []),
    )
    output_dir = Path(state["project"]["project_dir"]) / "outputs"
    state["reports"] = export_production_reports(state, output_dir)
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        _replace_section_entities(conn, run_id, sections)
    _save_run_state(run_id, state)
    return state


def _task_spec_prompt(
    state: dict[str, Any],
    blueprint: dict[str, Any],
    case_outlines: list[dict[str, Any]],
) -> str:
    payload = {
        "project_type": state.get("project", {}).get("project_type", ""),
        "tender_requirements": [
            {
                "id": item.get("id", ""),
                "category": item.get("category", ""),
                "requirement": compact_text(str(item.get("requirement", "")), 320),
            }
            for item in state.get("requirements", {}).get("requirements", [])[:100]
        ],
        "project_fact_names": list(
            dict.fromkeys(str(item.get("key", "")) for item in state.get("facts", []))
        ),
        "referenced_standards": [
            f"{item.get('code', '')} {item.get('title', '')}"
            for item in state.get("standards", [])[:80]
        ],
        "matched_case_outlines": case_outlines,
        "company_compact_start": [
            item.get("title", "") for item in blueprint.get("sections", [])
        ],
        "controlled_reference_modules": [
            item.get("title", "") for item in blueprint.get("reference_sections", [])
        ],
    }
    return (
        "请形成建议的一级章节任务书。一般建议5至12章，但必须依据资料决定，"
        "允许超出；合并或拆分都要保证每项招标要求有落点。历史案例只能借鉴结构，"
        "不得沿用其中项目参数。每章返回 title、purpose、required_inputs、components、acceptance。"
        "另返回 rationale，说明为何采用这个章节结构。JSON结构："
        '{"sections":[{"title":"","purpose":"","required_inputs":[],"components":[],"acceptance":[]}],'
        '"rationale":""}\n证据摘要：'
        + json.dumps(payload, ensure_ascii=False)
    )


def _case_outline_candidates(case_assets: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for asset in case_assets:
        if asset.get("asset_type") != "outline":
            continue
        content = asset.get("content")
        titles = [
            _clean_section_title(str(item.get("title", "")))
            for item in content if isinstance(item, dict)
            if int(item.get("level", 1) or 1) == 1
        ] if isinstance(content, list) else []
        titles = [title for title in titles if title]
        if titles:
            candidates.append(
                {
                    "case_id": asset.get("case_id"),
                    "case_title": asset.get("case_title", ""),
                    "titles": titles[:30],
                }
            )
    return candidates


def _closest_section_template(title: str, blueprint: dict[str, Any]) -> dict[str, Any]:
    candidates = [
        *blueprint.get("sections", []),
        *blueprint.get("reference_sections", []),
    ]
    if not candidates:
        return {}
    title_chars = set(re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", title))
    return max(
        candidates,
        key=lambda item: len(
            title_chars.intersection(
                set(re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", str(item.get("title", ""))))
            )
        ),
    )


def _clean_section_title(title: str) -> str:
    return re.sub(
        r"^\s*(?:第[一二三四五六七八九十百0-9]+章|[0-9]+(?:\.[0-9]+)*[、.．]?)\s*",
        "",
        title,
    ).strip()[:120]


def _clean_string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    result: list[str] = []
    for item in value:
        text = str(item).strip()
        if text and text not in result:
            result.append(text[:500])
    return result[:50]


def confirm_file_roles(run_id: int) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    state["file_roles_confirmed"] = True
    for item in state.get("files", []):
        item["role_confirmed"] = True
    if state.get("visual_jobs"):
        state["stage"] = "证据提取"
        state["stage_index"] = 3
        state["status"] = "visual_analysis_pending"
    else:
        state["stage"] = "编制任务书确认"
        state["stage_index"] = 6
        state["status"] = "awaiting_confirmation"
    state["file_roles_confirmed_at"] = _now()
    _save_run_state(run_id, state)
    return state


def mark_generation_started(run_id: int) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if state.get("status") == "generating":
        return state
    if state.get("status") not in {"ready_for_generation", "generation_failed", "partial_draft"}:
        raise ValueError("必须先确认《施组编制任务书》后才能生成正文")
    state["stage"] = "章节生成"
    state["stage_index"] = 7
    state["status"] = "generating"
    state["generation"] = {
        "total": len(state.get("sections", [])),
        "completed": 0,
        "successful": 0,
        "failed": 0,
        "started_at": _now(),
        "finished_at": "",
        "error": "",
    }
    _save_run_state(run_id, state)
    return state


def generate_production_docx(
    run_id: int,
    *,
    generate_missing: bool = True,
    run_text_review: bool = True,
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if state.get("status") in {"ready_for_generation", "generation_failed", "partial_draft"}:
        state = mark_generation_started(run_id)
    elif state.get("status") not in {"generating", "ready_for_assembly"}:
        raise ValueError("必须先确认《施组编制任务书》后才能生成正文")
    router = ModelRouter()
    text_model = next(item for item in router.status() if item["role"] == "text_master")
    generated_by_code: dict[str, dict[str, Any]] = {
        section["code"]: {"title": section["title"], "content": section["content"]}
        for section in state.get("sections", [])
        if section.get("status") == "generated" and section.get("content")
    }
    model_runs: list[dict[str, Any]] = list(state.get("model_runs", []))
    successful_sections = len(generated_by_code)
    sections = state.get("sections", [])
    pending_sections = [
        section for section in sections if section["code"] not in generated_by_code
    ]
    if pending_sections and not generate_missing:
        missing = "、".join(
            f"{section['code']} {section['title']}" for section in pending_sections[:5]
        )
        raise ValueError(f"尚未完成全部章节写作：{missing}")
    if pending_sections and not text_model["configured"]:
        raise RuntimeError("DEEPSEEK_API_KEY 未配置，系统不会用其他模型替代文字主脑")
    state["generation"].update(
        {
            "completed": successful_sections,
            "successful": successful_sections,
            "failed": 0,
        }
    )
    _save_run_state(run_id, state)
    workers = max(1, min(int(os.environ.get("PASS_BID_TEXT_WORKERS", "3")), 5))
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="deepseek-section") as executor:
        futures = {
            executor.submit(_generate_one_section, router, text_model, state, section): section
            for section in pending_sections
        }
        for future in as_completed(futures):
            section = futures[future]
            basis = section.get("project_basis", [])
            references = section.get("reference_basis", [])
            try:
                result = future.result()
                generated_by_code[section["code"]] = {
                    "title": section["title"],
                    "content": result["content"],
                }
                successful_sections += 1
                human_items = result["human_confirm_items"]
                section["status"] = "generated"
                section["completion"] = 100
                section["content"] = result["content"]
                for item in human_items:
                    state.setdefault("confirmations", []).append(
                        {
                            "category": "模型生成待确认",
                            "title": f"{section['title']}待确认",
                            "detail": item,
                            "severity": "medium",
                            "affected_sections": [section["title"]],
                            "status": "open",
                            "recommended_action": "补充项目资料或人工确认后重新生成本章。",
                        }
                    )
                model_runs.append(
                    {
                        "role": "text_master",
                        "provider": "deepseek",
                        "model": text_model["model"],
                        "task_type": "section_generation",
                        "section": section["title"],
                        "input_basis": basis + references,
                        "status": "success",
                    }
                )
                _persist_model_run(run_id, model_runs[-1])
            except Exception as exc:
                generated_by_code[section["code"]] = {
                    "title": section["title"],
                    "content": f"【待人工确认：本章模型生成失败，原因：{exc}】",
                }
                section["status"] = "blocked"
                section["content"] = generated_by_code[section["code"]]["content"]
                model_runs.append(
                    {
                        "role": "text_master",
                        "provider": "deepseek",
                        "model": text_model["model"],
                        "task_type": "section_generation",
                        "section": section["title"],
                        "input_basis": basis + references,
                        "status": "failed",
                        "error": str(exc),
                    }
                )
                _persist_model_run(run_id, model_runs[-1])
            state["model_runs"] = model_runs
            completed = len(generated_by_code)
            state["generation"].update(
                {
                    "completed": completed,
                    "successful": successful_sections,
                    "failed": completed - successful_sections,
                }
            )
            state["metrics"]["chapter_completion"] = round(
                sum(item["completion"] for item in sections) / max(len(sections), 1)
            )
            _save_run_state(run_id, state)
    if successful_sections == 0:
        state["model_runs"] = model_runs
        settings = get_settings()
        with db.db_session(settings.database_path) as conn:
            conn.execute(
                "UPDATE production_runs SET state_json = ?, updated_at = ? WHERE id = ?",
                (_json(state), _now(), run_id),
            )
        first_error = next(
            (item.get("error") for item in model_runs if item.get("status") == "failed"),
            "未知错误",
        )
        raise RuntimeError(f"DeepSeek未生成任何有效章节：{first_error}")
    generated_sections = [
        generated_by_code[section["code"]]
        for section in sections
    ]
    project_dir = Path(state["project"]["project_dir"])
    output_dir = project_dir / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    state["assets"] = generate_document_assets(state, output_dir)
    image_bindings = {
        "03": [
            ("施工总体流程图", state["assets"]["production_workflow_png"]),
            ("项目组织机构图（人员姓名和资格待确认）", state["assets"]["organization_chart_png"]),
        ],
        "05": [("施工进度计划结构图（具体日期与持续时间待确认）", state["assets"]["schedule_template_png"])],
        "07": [("质量控制流程图", state["assets"]["quality_flow_png"])],
        "08": [("安全管理流程图", state["assets"]["safety_flow_png"])],
    }
    drawing_output_dir = output_dir / "assets" / "drawings"
    for drawing in state.get("drawings", []):
        if drawing.get("status") != "confirmed" or not drawing.get("source_pdf_path"):
            continue
        try:
            image_path = render_confirmed_drawing(drawing, drawing_output_dir)
        except Exception as exc:
            state.setdefault("confirmations", []).append(
                {
                    "category": "图纸插入失败",
                    "title": drawing.get("title", "设计图纸"),
                    "detail": str(exc),
                    "severity": "high",
                    "affected_sections": drawing.get("applicable_sections", ["全局"]),
                    "status": "open",
                    "recommended_action": "检查图纸PDF页码和裁剪范围后重新装配。",
                }
            )
            continue
        targets = drawing.get("applicable_sections") or ["04"]
        for target in targets:
            code = str(target)
            if not re.fullmatch(r"\d{2}", code):
                matched = next(
                    (
                        section["code"]
                        for section in sections
                        if code in section.get("title", "")
                    ),
                    "04",
                )
                code = matched
            caption = (
                drawing.get("caption")
                or drawing.get("title")
                or drawing.get("drawing_no")
                or "设计图纸引用"
            )
            if drawing.get("placement") == "landscape_page":
                caption = f"{caption}|landscape"
            image_bindings.setdefault(code, []).append(
                (
                    caption,
                    str(image_path),
                )
            )
    for section, generated in zip(sections, generated_sections, strict=True):
        for caption, image_path in image_bindings.get(section["code"], []):
            generated["content"] += f"\n\n![{caption}]({image_path})"
    output_path = output_dir / f"{safe_filename(state['project']['name'])}_施工组织设计_初稿.docx"
    layout_profile = _recommended_layout_profile(state)
    result = generate_docx(
        title=f"{state['project']['name']} 施工组织设计",
        output_path=output_path,
        outline=[
            {"level": 1, "title": section["title"], "order": index}
            for index, section in enumerate(state["sections"], start=1)
        ],
        sections=generated_sections,
        requirements=state.get("requirements", {}),
        response_matrix=state.get("response_matrix", []),
        layout_profile=layout_profile,
        include_requirement_response=True,
    )
    state["model_runs"] = model_runs
    state["docx"] = result
    state["outputs"] = {
        **state.get("outputs", {}),
        "docx_path": result.get("docx_path", ""),
    }
    state["compliance"] = check_docx_compliance(
        docx_path=output_path,
        requirements=state.get("requirements", {}),
    )
    _apply_compliance_to_matrix(state)
    state["document_audit"] = _audit_generated_docx(output_path, state)
    _append_audit_confirmations(state)
    if run_text_review:
        state["text_review"] = _run_final_text_review(router, state)
    else:
        state["text_review"] = state.get("text_review") or {
            "status": "pending_text_review",
            "issues": [],
            "message": "等待DeepSeek完成跨章节文字复核。",
        }
    try:
        state["visual_report"] = visual_check_document(
            target_path=output_path,
            settings=get_settings(),
            reference_layout=layout_profile,
            max_pages=8,
        )
    except Exception as exc:
        state["visual_report"] = {"status": "error", "message": str(exc)}
    if state["visual_report"].get("status") not in {"ready", "disabled"}:
        state.setdefault("confirmations", []).append(
            {
                "category": "视觉复核待完成",
                "title": "DOCX视觉复核未取得模型结论",
                "detail": state["visual_report"].get("message", "视觉模型未配置或渲染失败"),
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "配置千问视觉模型后重新执行视觉复核，或人工检查导出的PDF。",
            }
        )
    state["stage"] = "输出交付"
    state["stage_index"] = 12
    state["status"] = "draft_generated" if successful_sections == len(state["sections"]) else "partial_draft"
    state["generation"]["finished_at"] = _now()
    state["metrics"]["chapter_completion"] = round(
        sum(section["completion"] for section in state["sections"]) / max(len(state["sections"]), 1)
    )
    state["metrics"]["open_confirmations"] = sum(
        1 for item in state["confirmations"] if item.get("status") == "open"
    )
    state["metrics"]["requirement_coverage"] = round(
        state["compliance"].get("covered_count", 0)
        / max(state["compliance"].get("requirement_count", 0), 1)
        * 100
    )
    state["reports"] = export_production_reports(state, output_dir)
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        conn.execute(
            """
            UPDATE production_runs SET stage = ?, status = ?, state_json = ?, updated_at = ?
            WHERE id = ?
            """,
            (state["stage"], state["status"], _json(state), _now(), run_id),
        )
    return state


def get_production_section_task(
    run_id: int,
    section_code: str = "",
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if state.get("status") in {"ready_for_generation", "generation_failed", "partial_draft"}:
        state = mark_generation_started(run_id)
    if state.get("status") != "generating":
        raise ValueError("必须先确认《施组编制任务书》后才能领取章节写作任务")

    pending = [
        section
        for section in state.get("sections", [])
        if section.get("status") != "generated" or not section.get("content")
    ]
    if not pending:
        return {
            "run_id": int(run_id),
            "done": True,
            "message": "全部章节已经人工确认，可以装配DOCX。",
            "completed": len(state.get("sections", [])),
            "total": len(state.get("sections", [])),
        }

    selected = None
    if section_code:
        selected = next(
            (item for item in pending if item.get("code") == section_code),
            None,
        )
        if selected is None:
            raise ValueError(f"待写章节不存在或已完成：{section_code}")
    else:
        selected = pending[0]
    return {
        "run_id": int(run_id),
        "done": False,
        "section_code": selected["code"],
        "title": selected["title"],
        "purpose": selected.get("purpose", ""),
        "acceptance": selected.get("acceptance", []),
        "components": selected.get("components", []),
        "project_basis": selected.get("project_basis", []),
        "reference_basis": selected.get("reference_basis", []),
        "missing_inputs": selected.get("missing_inputs", []),
        "writing_prompt": _section_prompt(state, selected),
        "completed": len(state.get("sections", [])) - len(pending),
        "total": len(state.get("sections", [])),
    }


def save_production_section(
    run_id: int,
    section_code: str,
    content: str,
    *,
    author_model: str = "DeepSeek",
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if state.get("status") in {"ready_for_generation", "generation_failed", "partial_draft"}:
        state = mark_generation_started(run_id)
    if state.get("status") != "generating":
        raise ValueError("当前生产任务不在章节生成阶段")
    section = next(
        (item for item in state.get("sections", []) if item.get("code") == section_code),
        None,
    )
    if section is None:
        raise ValueError(f"章节不存在：{section_code}")
    text = str(content or "").strip()
    if len(text) < 200:
        raise ValueError("章节正文过短；必须提交完整、可复核的章节内容")

    text = _strip_duplicate_section_heading(text, section.get("title", ""))
    text = _sanitize_generated_structure(text)
    text, findings = _sanitize_unsupported_claims(text, state, section)
    text = _append_basis_traceability(text, section)
    section["content"] = text
    section["status"] = "generated"
    section["completion"] = 100
    section["model"] = author_model
    for finding in findings:
        state.setdefault("confirmations", []).append(
            {
                "category": "模型生成待确认",
                "title": f"{section['title']}待确认",
                "detail": finding,
                "severity": "medium",
                "affected_sections": [section["title"]],
                "status": "open",
                "recommended_action": "补充项目资料或由人工核对该项依据。",
            }
        )

    generated = [
        item
        for item in state.get("sections", [])
        if item.get("status") == "generated" and item.get("content")
    ]
    total = len(state.get("sections", []))
    state.setdefault("generation", {}).update(
        {
            "total": total,
            "completed": len(generated),
            "successful": len(generated),
            "failed": 0,
        }
    )
    state.setdefault("metrics", {})["chapter_completion"] = round(
        len(generated) / max(total, 1) * 100
    )
    model_run = {
        "role": "text_master",
        "provider": "deepseek",
        "model": author_model,
        "task_type": "section_generation",
        "section": section["title"],
        "input_basis": section.get("project_basis", [])
        + section.get("reference_basis", []),
        "status": "success",
    }
    state.setdefault("model_runs", []).append(model_run)
    _persist_model_run(run_id, model_run)
    _save_run_state(run_id, state)

    next_section = next(
        (
            item
            for item in state.get("sections", [])
            if item.get("status") != "generated" or not item.get("content")
        ),
        None,
    )
    return {
        "run_id": int(run_id),
        "saved": section_code,
        "completed": len(generated),
        "total": total,
        "done": next_section is None,
        "next_section": (
            {"section_code": next_section["code"], "title": next_section["title"]}
            if next_section
            else None
        ),
        "human_confirmation_count": len(findings),
    }


def generate_production_section(
    run_id: int,
    section_code: str,
    *,
    revision_instruction: str = "",
) -> dict[str, Any]:
    """Generate or revise exactly one chapter and wait for human approval."""
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    if state.get("status") in {"ready_for_generation", "generation_failed", "partial_draft"}:
        state = mark_generation_started(run_id)
    if state.get("status") not in {"generating", "ready_for_assembly"}:
        raise ValueError("必须先确认《施组编制任务书》后才能生成章节")
    sections = state.get("sections", [])
    section = next((item for item in sections if item.get("code") == section_code), None)
    if section is None:
        raise ValueError(f"章节不存在：{section_code}")
    first_unapproved = next(
        (item for item in sections if item.get("status") != "generated"),
        None,
    )
    if first_unapproved and first_unapproved.get("code") != section_code:
        raise ValueError(
            f"请先确认 {first_unapproved['code']} {first_unapproved['title']}"
        )

    router = ModelRouter()
    text_model = next(item for item in router.status() if item["role"] == "text_master")
    if not text_model["configured"]:
        raise RuntimeError("DEEPSEEK_API_KEY 未配置，无法生成章节正文")
    result = _generate_one_section(
        router,
        text_model,
        state,
        section,
        revision_instruction=revision_instruction,
    )
    previous_content = str(section.get("content", ""))
    section.setdefault("revisions", []).append(
        {
            "created_at": _now(),
            "instruction": revision_instruction,
            "previous_content": previous_content,
            "model": result["model"],
        }
    )
    section["content"] = result["content"]
    section["status"] = "awaiting_approval"
    section["completion"] = 90
    section["model"] = result["model"]
    for item in result["human_confirm_items"]:
        state.setdefault("confirmations", []).append(
            {
                "category": "章节生成待确认",
                "title": f"{section['title']}待确认",
                "detail": str(item),
                "severity": "medium",
                "affected_sections": [section["title"]],
                "status": "open",
                "recommended_action": "补充项目依据，或在确认章节前核对该内容。",
            }
        )
    model_run = {
        "role": "text_master",
        "provider": "deepseek",
        "model": result["model"],
        "task_type": "section_revision" if revision_instruction else "section_generation",
        "section": section["title"],
        "input_basis": section.get("project_basis", [])
        + section.get("reference_basis", []),
        "status": "success",
    }
    state.setdefault("model_runs", []).append(model_run)
    _persist_model_run(run_id, model_run)
    state["status"] = "generating"
    state.setdefault("generation", {})["current_section"] = section_code
    state["generation"]["awaiting_approval"] = section_code
    state.setdefault("metrics", {})["chapter_completion"] = round(
        sum(item.get("completion", 0) for item in sections) / max(len(sections), 1)
    )
    _save_run_state(run_id, state)
    return state


def approve_production_section(run_id: int, section_code: str) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    section = next(
        (item for item in state.get("sections", []) if item.get("code") == section_code),
        None,
    )
    if section is None:
        raise ValueError(f"章节不存在：{section_code}")
    if section.get("status") != "awaiting_approval" or not section.get("content"):
        raise ValueError("该章节尚无待确认正文")
    section["status"] = "generated"
    section["completion"] = 100
    section["approved_at"] = _now()
    approved = [
        item
        for item in state.get("sections", [])
        if item.get("status") == "generated"
    ]
    total = len(state.get("sections", []))
    state.setdefault("generation", {}).update(
        {
            "total": total,
            "completed": len(approved),
            "successful": len(approved),
            "failed": 0,
            "awaiting_approval": "",
        }
    )
    state["metrics"]["chapter_completion"] = round(len(approved) / max(total, 1) * 100)
    state["status"] = "ready_for_assembly" if len(approved) == total else "generating"
    state["stage"] = "文档装配" if len(approved) == total else "章节生成"
    state["stage_index"] = 11 if len(approved) == total else 7
    _save_run_state(run_id, state)
    return state


def assemble_production_docx(run_id: int) -> dict[str, Any]:
    """Assemble only after every AI-generated section is human-approved."""
    state = generate_production_docx(
        run_id,
        generate_missing=False,
        run_text_review=False,
    )
    docx_path = Path(str(state.get("docx", {}).get("docx_path", "")))
    pdf_path = docx_path.with_suffix(".pdf")
    state["pdf"] = export_docx_to_pdf(docx_path, pdf_path)
    if state["pdf"].get("status") == "ready":
        state.setdefault("outputs", {})["pdf_path"] = str(pdf_path)
    if state["pdf"].get("status") != "ready":
        state.setdefault("confirmations", []).append(
            {
                "category": "PDF导出待完成",
                "title": "DOCX已生成，但PDF转换未完成",
                "detail": state["pdf"].get("message", "Word或LibreOffice转换失败"),
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "检查Microsoft Word后重新装配，或人工另存为PDF。",
            }
        )
    _save_run_state(run_id, state)
    return state


def record_production_text_review(
    run_id: int,
    *,
    status: str,
    issues: list[dict[str, Any]] | None = None,
    summary: str = "",
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    normalized_status = str(status or "").strip().lower()
    if normalized_status not in {"passed", "needs_review", "blocked"}:
        raise ValueError("status必须为 passed、needs_review 或 blocked")
    accepted, dismissed = _validate_review_issues(state, issues or [])
    state["text_review"] = {
        "status": normalized_status if accepted or normalized_status == "passed" else "passed",
        "issues": accepted,
        "dismissed_issues": dismissed,
        "summary": str(summary or "").strip(),
        "reviewer": "DeepSeek",
        "reviewed_at": _now(),
    }
    for issue in accepted:
        state.setdefault("confirmations", []).append(
            {
                "category": "DeepSeek文字复核",
                "title": issue.get("title", "跨章节文字问题"),
                "detail": issue.get("detail", ""),
                "severity": issue.get("severity", "medium"),
                "affected_sections": issue.get("affected_sections", ["全局"]),
                "status": "open",
                "recommended_action": issue.get(
                    "recommended_action",
                    "由DeepSeek修订相关章节或人工确认。",
                ),
            }
        )
    _save_run_state(run_id, state)
    return {
        "run_id": int(run_id),
        "text_review": state["text_review"],
        "status": state.get("status"),
    }


def mark_generation_failed(run_id: int, error: str) -> dict[str, Any] | None:
    state = get_run_state(run_id)
    if state is None:
        return None
    state["status"] = "generation_failed"
    state.setdefault("generation", {})["error"] = error
    state["generation"]["finished_at"] = _now()
    _save_run_state(run_id, state)
    return state


def review_production_draft(
    run_id: int,
    *,
    run_text_review: bool = True,
) -> dict[str, Any]:
    state = get_run_state(run_id)
    if state is None:
        raise ValueError(f"production run not found: {run_id}")
    docx_path = Path(str(state.get("docx", {}).get("docx_path", "")))
    if not docx_path.exists():
        raise ValueError("当前生产任务没有可复核的DOCX")
    state["confirmations"] = [
        item
        for item in state.get("confirmations", [])
        if item.get("category") not in {"最终文字复核", "视觉复核待完成", "文档装配复核"}
    ]
    state["compliance"] = check_docx_compliance(
        docx_path=docx_path,
        requirements=state.get("requirements", {}),
    )
    _apply_compliance_to_matrix(state)
    state["document_audit"] = _audit_generated_docx(docx_path, state)
    _append_audit_confirmations(state)
    if run_text_review:
        router = ModelRouter()
        state["text_review"] = _run_final_text_review(router, state)
    layout_profile = _recommended_layout_profile(state)
    try:
        state["visual_report"] = visual_check_document(
            target_path=docx_path,
            settings=get_settings(),
            reference_layout=layout_profile,
            max_pages=8,
        )
    except Exception as exc:
        state["visual_report"] = {"status": "error", "message": str(exc)}
    if state["visual_report"].get("status") not in {"ready", "disabled"}:
        state.setdefault("confirmations", []).append(
            {
                "category": "视觉复核待完成",
                "title": "DOCX视觉复核未取得模型结论",
                "detail": state["visual_report"].get("message", "视觉模型未配置或渲染失败"),
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "配置千问视觉模型后重新执行视觉复核，或人工检查导出的PDF。",
            }
        )
    state["metrics"]["requirement_coverage"] = round(
        state["compliance"].get("covered_count", 0)
        / max(state["compliance"].get("requirement_count", 0), 1)
        * 100
    )
    state["metrics"]["open_confirmations"] = sum(
        1 for item in state.get("confirmations", []) if item.get("status") == "open"
    )
    state["metrics"]["high_risks"] = sum(
        1
        for item in state.get("confirmations", [])
        if item.get("status") == "open" and item.get("severity") == "high"
    )
    state["reports"] = export_production_reports(state, docx_path.parent)
    state["reviewed_at"] = _now()
    _save_run_state(run_id, state)
    return state


def _generate_one_section(
    router: ModelRouter,
    text_model: dict[str, Any],
    state: dict[str, Any],
    section: dict[str, Any],
    revision_instruction: str = "",
) -> dict[str, Any]:
    system = (
        "你是水利工程施工组织设计文字主脑。只允许使用输入中的项目依据和编制依据。"
        "不得编造日期、工程量、人员、设备、标准条款或设计参数。"
        "缺少信息时使用【待人工确认：...】。必须返回JSON对象，字段为content和human_confirm_items。"
    )
    result: dict[str, Any] | None = None
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            prompt = _section_prompt(state, section)
            if revision_instruction:
                prompt += (
                    "\n\n用户修改要求："
                    + revision_instruction
                    + "\n请在保留所有有效依据和禁止虚构约束的前提下重写本章。"
                )
            result = router.complete_json(
                role="text_master",
                system=system + (" 上次输出无法解析，请只输出合法JSON。" if attempt else ""),
                prompt=prompt,
                timeout=300,
                max_tokens=int(os.environ.get("PASS_BID_SECTION_MAX_TOKENS", "3200")),
            )
            break
        except RuntimeError as exc:
            last_error = exc
            if "invalid JSON" not in str(exc) or attempt == 2:
                raise
    if result is None:
        raise RuntimeError(str(last_error or "DeepSeek未返回结果"))
    content = str(result.get("content", "")).strip()
    if not content:
        raise RuntimeError("DeepSeek返回空正文")
    content = _strip_duplicate_section_heading(content, section["title"])
    content = _sanitize_generated_structure(content)
    content, audit_items = _sanitize_unsupported_claims(content, state, section)
    content = _append_basis_traceability(content, section)
    return {
        "content": content,
        "human_confirm_items": [
            *[str(item) for item in result.get("human_confirm_items", [])],
            *audit_items,
        ],
        "model": text_model["model"],
    }


def _strip_duplicate_section_heading(content: str, section_title: str) -> str:
    lines = content.splitlines()
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.sub(r"^#{1,4}\s*", "", stripped)
        heading = re.sub(r"^第[一二三四五六七八九十百\d]+章\s*", "", heading)
        if heading.strip(" ：:") == section_title.strip(" ：:"):
            del lines[index]
        break
    return "\n".join(lines).strip()


def _sanitize_generated_structure(content: str) -> str:
    """Keep model output as prose; controlled diagrams and basis tables are assembled separately."""
    had_mermaid = bool(re.search(r"```mermaid\b", content, flags=re.IGNORECASE))
    content = re.sub(
        r"```mermaid\b.*?```",
        "",
        content,
        flags=re.IGNORECASE | re.DOTALL,
    )
    content = re.sub(r"```(?:\w+)?\s*.*?```", "", content, flags=re.DOTALL)
    content = re.sub(r"</?(?:sub|sup|br)\s*/?>", "", content, flags=re.IGNORECASE)
    content = re.sub(
        r"(?ms)^#{1,4}\s*本章依据追溯\s*$.*\Z",
        "",
        content,
    )
    content = _strip_ascii_diagram_subsections(content)
    if had_mermaid:
        content = content.rstrip() + "\n\n图示由系统以原生图形组件生成，见本节末。"
    return re.sub(r"\n{3,}", "\n\n", content).strip()


def _strip_ascii_diagram_subsections(content: str) -> str:
    lines = content.splitlines()
    output: list[str] = []
    index = 0
    heading_pattern = re.compile(r"^(#{1,4})\s+.*(?:组织机构图|流程图)\s*$")
    while index < len(lines):
        match = heading_pattern.match(lines[index].strip())
        if not match:
            output.append(lines[index])
            index += 1
            continue
        level = len(match.group(1))
        end = index + 1
        while end < len(lines):
            next_heading = re.match(r"^(#{1,4})\s+", lines[end].strip())
            if next_heading and len(next_heading.group(1)) <= level:
                break
            end += 1
        segment = lines[index + 1:end]
        diagram_lines = sum(
            1
            for line in segment
            if re.search(r"[|│─—┌┐└┘├┤┬┴┼↓↑→←]", line)
            or line.strip() in {"...", "…"}
        )
        output.append(lines[index])
        if diagram_lines >= 3:
            output.extend(["", "图示由系统以原生图形组件生成，见本节末。", ""])
        else:
            output.extend(segment)
        index = end
    return "\n".join(output)


def _append_basis_traceability(content: str, section: dict[str, Any]) -> str:
    rows = ["### 本章依据追溯", "", "| 依据类型 | 来源 |", "|---|---|"]
    for item in section.get("project_basis", [])[:12]:
        rows.append(f"| 项目依据 | {str(item).replace('|', '／')} |")
    for item in section.get("reference_basis", [])[:12]:
        rows.append(f"| 编制依据 | {str(item).replace('|', '／')} |")
    if section.get("missing_inputs"):
        rows.append(
            "| 待确认输入 | "
            + "、".join(str(item).replace("|", "／") for item in section["missing_inputs"])
            + " |"
        )
    return content.rstrip() + "\n\n" + "\n".join(rows)


def _persist_model_run(run_id: int, item: dict[str, Any]) -> None:
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        cursor = conn.execute(
            """
            INSERT INTO model_runs
            (production_run_id, role, provider, model, task_type, input_basis_json,
             result_json, status, error, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                item.get("role", ""),
                item.get("provider", ""),
                item.get("model", ""),
                item.get("task_type", ""),
                _json(item.get("input_basis", [])),
                _json(item.get("result", {})),
                item.get("status", ""),
                item.get("error", ""),
                _now(),
            ),
        )
        item["db_id"] = int(cursor.lastrowid)


def _recommended_layout_profile(state: dict[str, Any]) -> dict[str, Any] | None:
    layout_id = next(
        (
            asset.get("layout_profile_id")
            for asset in state.get("case_assets", [])
            if asset.get("asset_type") == "layout" and asset.get("layout_profile_id")
        ),
        None,
    )
    if not layout_id:
        return None
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        record = db.get_layout_profile(conn, int(layout_id))
    return record.get("profile") if record else None


def _apply_compliance_to_matrix(state: dict[str, Any]) -> None:
    missing_ids = {
        item.get("requirement_id")
        for item in state.get("compliance", {}).get("missing_requirements", [])
    }
    for row in state.get("response_matrix", []):
        row["draft_status"] = (
            "missing"
            if row.get("requirement_id") in missing_ids
            else "covered"
        )


def _audit_generated_docx(docx_path: Path, state: dict[str, Any]) -> dict[str, Any]:
    from docx import Document

    document = Document(str(docx_path))
    heading_texts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name.startswith("Heading")
    ]
    missing_sections = [
        section["title"]
        for section in state.get("sections", [])
        if not any(section["title"] in heading for heading in heading_texts)
    ]
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    placeholder_count = len(
        re.findall(r"【[^】]*(?:待人工确认|待完善|待确认)[^】]*】", body_text)
    )
    result = {
        "status": "passed",
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "image_count": len(document.inline_shapes),
        "missing_sections": missing_sections,
        "placeholder_count": placeholder_count,
        "checks": {
            "all_sections_present": not missing_sections,
            "has_word_tables": len(document.tables) > 0,
            "has_generated_figures": len(document.inline_shapes) >= 5,
            "field_update_ready": state.get("docx", {}).get("field_update", {}).get("status") == "ready",
        },
    }
    failed_checks = [name for name, passed in result["checks"].items() if not passed]
    if failed_checks:
        result["status"] = "needs_review"
        result["failed_checks"] = failed_checks
    return result


def _append_audit_confirmations(state: dict[str, Any]) -> None:
    audit = state.get("document_audit", {})
    if audit.get("status") == "passed":
        return
    state.setdefault("confirmations", []).append(
        {
            "category": "文档装配复核",
            "title": "DOCX结构验收存在未通过项",
            "detail": "、".join(audit.get("failed_checks", [])),
            "severity": "high",
            "affected_sections": audit.get("missing_sections", []) or ["全局"],
            "status": "open",
            "recommended_action": "修复装配或内容组件后重新生成DOCX。",
        }
    )


def _run_final_text_review(
    router: ModelRouter,
    state: dict[str, Any],
) -> dict[str, Any]:
    if os.environ.get("PASS_BID_FINAL_TEXT_REVIEW", "true").lower() in {"0", "false", "off"}:
        return {"status": "disabled", "issues": []}
    sections = [
        {
            "title": section["title"],
            "content_excerpt": str(section.get("content", ""))[:1200],
            "project_basis": section.get("project_basis", [])[:6],
            "reference_basis": section.get("reference_basis", [])[:6],
        }
        for section in state.get("sections", [])
    ]
    try:
        review: dict[str, Any] | None = None
        for attempt in range(3):
            try:
                review = router.complete_json(
                    role="text_master",
                    system=(
                        "你是施工组织设计最终文字复核员。只找问题，不重写正文。"
                        "检查招标要求漏项、跨章数字矛盾、无依据规范条款、人员机械进度不一致。"
                        "必须返回合法JSON对象，issues最多10项，每项detail不超过100字。"
                        + (" 上次JSON无法解析，请缩短内容并严格闭合括号。" if attempt else "")
                    ),
                    prompt=_json(
                        {
                            "requirements": state.get("requirements", {}).get("requirements", []),
                            "facts": state.get("facts", []),
                            "standard_clauses": state.get("standard_clauses", []),
                            "sections": sections,
                            "generated_assets": state.get("assets", {}),
                            "document_audit": state.get("document_audit", {}),
                            "判断规则": [
                                "generated_assets或document_audit已证明存在的图表，不得判为缺失",
                                "缺少项目参数但已明确标注待人工确认，应报告为交付前输入缺口，不得称为模型虚构",
                            ],
                            "输出": {
                                "status": "passed或needs_review",
                                "issues": [
                                    {
                                        "title": "问题",
                                        "detail": "说明",
                                        "severity": "low/medium/high",
                                        "affected_sections": ["章节"],
                                    }
                                ],
                            },
                        }
                    ),
                    timeout=300,
                    max_tokens=2200,
                )
                break
            except RuntimeError as exc:
                if "invalid JSON" not in str(exc) or attempt == 2:
                    raise
        if review is None:
            raise RuntimeError("DeepSeek最终复核未返回结果")
        accepted_issues, dismissed_issues = _validate_review_issues(
            state,
            review.get("issues", [])[:100],
        )
        review["issues"] = accepted_issues
        review["dismissed_issues"] = dismissed_issues
        for issue in accepted_issues:
            state.setdefault("confirmations", []).append(
                {
                    "category": "最终文字复核",
                    "title": str(issue.get("title", "文字一致性问题")),
                    "detail": str(issue.get("detail", "")),
                    "severity": str(issue.get("severity", "medium"))
                    if str(issue.get("severity", "medium")) in {"low", "medium", "high"}
                    else "medium",
                    "affected_sections": issue.get("affected_sections", []) or ["全局"],
                    "status": "open",
                    "recommended_action": "对照项目证据和招标要求修订后重新复核。",
                }
            )
        return {"status": review.get("status", "ready"), **review}
    except Exception as exc:
        state.setdefault("confirmations", []).append(
            {
                "category": "最终文字复核",
                "title": "DeepSeek最终文字复核未完成",
                "detail": str(exc),
                "severity": "medium",
                "affected_sections": ["全局"],
                "status": "open",
                "recommended_action": "恢复模型服务后重新运行复核。",
            }
        )
        return {"status": "error", "issues": [], "error": str(exc)}


def _validate_review_issues(
    state: dict[str, Any],
    issues: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Reject review claims contradicted by the assembled evidence; retain the audit trail."""
    all_text = "\n".join(str(item.get("content", "")) for item in state.get("sections", []))
    fact_text = _json(state.get("facts", []))
    asset_keys = set(state.get("assets", {}))
    verified_clauses = {
        str(item.get("clause_id", ""))
        for item in state.get("standard_clauses", [])
        if item.get("verification_status") == "verified_local_source"
    }
    section_titles = {str(item.get("title", "")) for item in state.get("sections", [])}
    accepted: list[dict[str, Any]] = []
    dismissed: list[dict[str, Any]] = []
    for issue in issues:
        claim = f"{issue.get('title', '')} {issue.get('detail', '')}"
        reason = ""
        numeric_claims = re.findall(
            r"\d+(?:\.\d+)?\s*(?:日历天|天|小时|分钟|mm|cm|m|km|%|％)",
            claim,
            flags=re.IGNORECASE,
        )
        absent_claims = [
            token
            for token in numeric_claims
            if re.sub(r"\s+", "", token) not in re.sub(r"\s+", "", all_text + fact_text)
        ]
        if absent_claims:
            reason = f"正文及项目事实中不存在被指称数字：{'、'.join(absent_claims)}"
        elif any(word in claim for word in ("未独立成章", "未专设章节", "缺少章节")) and any(
            title in section_titles for title in issue.get("affected_sections", [])
        ):
            reason = "章节清单和DOCX结构审计已确认对应独立章节存在"
        elif (
            any(word in claim for word in ("进度图", "横道图", "网络图"))
            and "schedule_template_png" in asset_keys
        ):
            reason = "资产清单和DOCX结构审计已确认进度图存在"
        elif (
            "组织机构图" in claim
            and "organization_chart_png" in asset_keys
        ):
            reason = "资产清单和DOCX结构审计已确认组织机构图存在"
        elif (
            "流程图" in claim
            and {"quality_flow_png", "safety_flow_png"}.intersection(asset_keys)
        ):
            reason = "资产清单和DOCX结构审计已确认流程图存在"
        elif any(
            word in claim for word in ("无法验证", "未验证", "无依据", "需验证", "原文支撑")
        ):
            cited = set(re.findall(r"(\d+(?:\.\d+){1,3})\s*条", claim))
            if cited and cited.issubset(verified_clauses):
                reason = "引用条款已由本地规范原文解析并标记为已验证"
        if reason:
            dismissed.append({**issue, "dismissed_reason": reason})
        else:
            accepted.append(issue)
    return accepted, dismissed


def _sanitize_unsupported_claims(
    content: str,
    state: dict[str, Any],
    section: dict[str, Any],
) -> tuple[str, list[str]]:
    """Block unsupported project quantities and clause citations before DOCX assembly."""
    allowed_text = _json(
        {
            "facts": state.get("facts", []),
            "requirements": section.get("requirements", []),
            "project_basis": section.get("project_basis", []),
            "standard_clauses": state.get("standard_clauses", []),
        }
    ).replace(" ", "")
    findings: list[str] = []
    seen: set[str] = set()
    verified_clause_tokens = {
        re.sub(r"\s+", "", f"{clause.get('standard_code', '')}第{clause.get('clause_id', '')}条")
        for clause in state.get("standard_clauses", [])
        if clause.get("verification_status") == "verified_local_source"
    }
    allowed_quantity_tokens = {
        re.sub(r"\s+", "", f"{fact.get('value', '')}{fact.get('unit', '')}")
        .replace("日历天", "天")
        .replace("％", "%")
        for fact in state.get("facts", [])
        if str(fact.get("value", "")).strip()
    }
    quantity_pattern = re.compile(
        r"(?<![A-Za-z0-9_.-])\d+(?:\.\d+)?\s*(?:日历天|天|小时|分钟|秒|个月|年|%|％|"
        r"mm|cm|m|km|m²|m2|㎡|m³|m3|立方米|平方米|"
        r"kW|kw|MPa|kPa|kN|N|dB|℃|°C|台|套|辆|人|班|次)(?![A-Za-z0-9_])",
        re.IGNORECASE,
    )
    clause_pattern = re.compile(
        r"(?:(?:GB|GB/T|SL|SL/T|DL/T|JGJ|CJJ|JTG|NB/T|DB\d{2}/T?)\s*"
        r"\d+(?:\.\d+)?(?:[-—–]\d{4})?)\s*(?:第\s*\d+(?:\.\d+)*\s*条|§\s*\d+(?:\.\d+)*)",
        re.IGNORECASE,
    )

    def replace_claim(match: re.Match[str]) -> str:
        token = re.sub(r"\s+", "", match.group(0))
        if any(verified in token for verified in verified_clause_tokens):
            return match.group(0)
        normalized_quantity = token.replace("日历天", "天").replace("％", "%")
        if (
            token in allowed_text
            or normalized_quantity in allowed_quantity_tokens
            or "待人工确认" in content[max(0, match.start() - 12):match.end() + 12]
        ):
            return match.group(0)
        if token not in seen:
            seen.add(token)
            findings.append(f"检测到无证据数字或条款引用“{match.group(0)}”，已在正文中阻断。")
        return f"【待人工确认：{match.group(0)}的项目或规范依据】"

    def restore_supported_placeholder(match: re.Match[str]) -> str:
        candidate = match.group(1)
        normalized = (
            re.sub(r"\s+", "", candidate)
            .replace("日历天", "天")
            .replace("％", "%")
        )
        if normalized in allowed_quantity_tokens or any(
            verified in re.sub(r"\s+", "", candidate)
            for verified in verified_clause_tokens
        ):
            return candidate
        return match.group(0)

    content = re.sub(
        r"【待人工确认：(.+?)的项目或规范依据】",
        restore_supported_placeholder,
        content,
    )
    sanitized = clause_pattern.sub(replace_claim, content)
    sanitized = quantity_pattern.sub(replace_claim, sanitized)
    sanitized = _sanitize_numeric_table_cells(
        sanitized,
        state=state,
        findings=findings,
    )
    return sanitized, findings


def _sanitize_numeric_table_cells(
    content: str,
    *,
    state: dict[str, Any],
    findings: list[str],
) -> str:
    """Block bare model-generated quantities whose units only appear in table headers."""
    lines = content.splitlines()
    allowed_values = {
        re.sub(r"\s+", "", f"{fact.get('value', '')}")
        for fact in state.get("facts", [])
        if str(fact.get("value", "")).strip()
    }
    quantity_headers = (
        "数量",
        "人数",
        "工程量",
        "日历天数",
        "天数",
        "工期",
        "持续时间",
        "平均人数",
        "高峰人数",
        "储备数量",
        "计划进场",
    )
    index = 0
    while index + 1 < len(lines):
        if not (
            lines[index].strip().startswith("|")
            and lines[index + 1].strip().startswith("|")
            and re.fullmatch(r"[\s|:\-]+", lines[index + 1].strip())
        ):
            index += 1
            continue
        header = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        protected_columns = {
            column
            for column, title in enumerate(header)
            if any(keyword in title for keyword in quantity_headers)
        }
        row_index = index + 2
        while row_index < len(lines) and lines[row_index].strip().startswith("|"):
            cells = [cell.strip() for cell in lines[row_index].strip().strip("|").split("|")]
            changed = False
            for column in protected_columns:
                if column >= len(cells):
                    continue
                value = cells[column]
                normalized = re.sub(r"\s+", "", value)
                if (
                    re.fullmatch(r"\d+(?:\.\d+)?(?:[-~～]\d+(?:\.\d+)?)?", normalized)
                    and normalized not in allowed_values
                    and "待人工确认" not in value
                ):
                    cells[column] = f"【待人工确认：{value}】"
                    findings.append(f"检测到表格中的无证据数量“{value}”，已在正文中阻断。")
                    changed = True
            if changed:
                lines[row_index] = "| " + " | ".join(cells) + " |"
            row_index += 1
        index = row_index
    return "\n".join(lines)


def _save_run_state(run_id: int, state: dict[str, Any]) -> None:
    for section in state.get("sections", []):
        section["project_basis"] = list(dict.fromkeys(section.get("project_basis", [])))
        section["reference_basis"] = list(dict.fromkeys(section.get("reference_basis", [])))
        section["missing_inputs"] = list(dict.fromkeys(section.get("missing_inputs", [])))
    unique_confirmations: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for item in state.get("confirmations", []):
        key = (
            str(item.get("category", "")),
            str(item.get("title", "")),
            str(item.get("detail", "")),
            _json(item.get("affected_sections", [])),
        )
        existing = unique_confirmations.get(key)
        if existing is None or (
            existing.get("status") == "open" and item.get("status") != "open"
        ):
            unique_confirmations[key] = item
    state["confirmations"] = list(unique_confirmations.values())
    metrics = state.setdefault("metrics", {})
    standards = state.get("standards", [])
    ready_standards = sum(
        1 for item in standards if item.get("status") in {"available", "not_applicable"}
    )
    metrics["standards_count"] = len(standards)
    metrics["standards_ready"] = ready_standards
    metrics["standards_readiness"] = (
        round(ready_standards / len(standards) * 100) if standards else 100
    )
    metrics["chapter_completion"] = (
        round(
            sum(int(item.get("completion", 0)) for item in state.get("sections", []))
            / len(state.get("sections", []))
        )
        if state.get("sections")
        else 0
    )
    metrics["open_confirmations"] = sum(
        1 for item in state.get("confirmations", []) if item.get("status") == "open"
    )
    metrics["high_risks"] = sum(
        1
        for item in state.get("confirmations", [])
        if item.get("status") == "open" and item.get("severity") == "high"
    )
    settings = get_settings()
    with db.db_session(settings.database_path) as conn:
        conn.execute(
            """
            UPDATE production_runs SET stage = ?, status = ?, state_json = ?, updated_at = ?
            WHERE id = ?
            """,
            (state["stage"], state["status"], _json(state), _now(), run_id),
        )
        for section in state.get("sections", []):
            conn.execute(
                """
                UPDATE section_tasks
                SET status = ?, completion = ?, basis_json = ?, content_text = ?, model = ?, updated_at = ?
                WHERE production_run_id = ? AND section_code = ?
                """,
                (
                    section.get("status", ""),
                    int(section.get("completion", 0)),
                    _json(
                        {
                            "project": section.get("project_basis", []),
                            "reference": section.get("reference_basis", []),
                        }
                    ),
                    section.get("content", ""),
                    section.get("model", ""),
                    _now(),
                    run_id,
                    section.get("code", ""),
                ),
            )
            conn.execute(
                """
                UPDATE content_components
                SET status = ?, updated_at = ?
                WHERE production_run_id = ? AND section_code = ?
                """,
                (
                    "generated" if section.get("status") == "generated" else "not_started",
                    _now(),
                    run_id,
                    section.get("code", ""),
                ),
            )
        conn.execute("DELETE FROM project_facts WHERE production_run_id = ?", (run_id,))
        for fact in state.get("facts", []):
            conn.execute(
                """
                INSERT INTO project_facts
                (production_run_id, fact_key, value, unit, source_path, source_page,
                 source_excerpt, confidence, status, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id,
                    fact.get("key", ""),
                    fact.get("value", ""),
                    fact.get("unit", ""),
                    fact.get("source_path", ""),
                    fact.get("source_page"),
                    fact.get("source_excerpt", ""),
                    float(fact.get("confidence", 0)),
                    fact.get("status", "extracted"),
                    _now(),
                ),
            )
        conn.execute("DELETE FROM confirmation_items WHERE production_run_id = ?", (run_id,))
        for item in state.get("confirmations", []):
            conn.execute(
                """
                INSERT INTO confirmation_items
                (production_run_id, category, title, detail, severity, affected_sections_json,
                 status, resolution, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id,
                    item.get("category", ""),
                    item.get("title", ""),
                    item.get("detail", ""),
                    item.get("severity", "medium"),
                    _json(item.get("affected_sections", [])),
                    item.get("status", "open"),
                    item.get("resolution", ""),
                    _now(),
                    _now(),
                ),
            )
        conn.execute("DELETE FROM drawing_assets WHERE production_run_id = ?", (run_id,))
        for drawing in state.get("drawings", []):
            conn.execute(
                """
                INSERT INTO drawing_assets
                (production_run_id, source_dwg_path, source_pdf_path, source_page, drawing_no,
                 title, caption, crop_json, applicable_sections_json, placement, preview_path,
                 confidence, status, confirmed_at, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id,
                    drawing.get("source_dwg_path", ""),
                    drawing.get("source_pdf_path", ""),
                    drawing.get("source_page"),
                    drawing.get("drawing_no", ""),
                    drawing.get("title", ""),
                    drawing.get("caption", ""),
                    _json(drawing.get("crop", {})),
                    _json(drawing.get("applicable_sections", [])),
                    drawing.get("placement", "inline"),
                    drawing.get("preview_path", ""),
                    float(drawing.get("confidence", 0)),
                    drawing.get("status", "pending"),
                    drawing.get("confirmed_at", ""),
                    _now(),
                    _now(),
                ),
            )


def _section_prompt(state: dict[str, Any], section: dict[str, Any]) -> str:
    case_asset_ids = set(section.get("case_asset_ids", []))
    case_assets = [
        {
            "title": asset.get("title"),
            "reuse_rule": asset.get("reuse_rule"),
            "content": asset.get("content"),
            "source_path": asset.get("source_path"),
        }
        for asset in state.get("case_assets", [])
        if asset.get("asset_id") in case_asset_ids
    ][:12]
    standard_clauses = [
        clause
        for clause in state.get("standard_clauses", [])
        if section.get("title") in clause.get("applicable_sections", [])
    ][:30]
    return _json(
        {
            "任务": f"编写《{section['title']}》完整初稿",
            "章节目的": section["purpose"],
            "必须包含的组件": section["components"],
            "完成标准": section["acceptance"],
            "项目依据": section["project_basis"],
            "编制依据": section["reference_basis"],
            "已验证规范条款": standard_clauses,
            "历史案例资产": case_assets,
            "对应招标要求": section["requirements"],
            "缺失输入": section["missing_inputs"],
            "全局项目事实": state.get("facts", [])[:80],
            "约束": [
                "仅依据给定资料写作",
                "无来源的具体参数必须标成人工确认项",
                "历史案例只允许复用结构和表达，不允许复用参数",
                "规范性结论必须引用输入中的已验证规范条款，只有规范名称时不得编造条款号",
                "使用水利施工现场语言，避免空泛口号",
                "正文使用Markdown，可包含表格",
            ],
            "输出JSON": {"content": "Markdown正文", "human_confirm_items": ["待确认事项"]},
        }
    )
