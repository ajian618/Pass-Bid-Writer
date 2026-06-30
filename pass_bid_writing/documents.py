from __future__ import annotations

import os
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


COLOR_KEYWORDS = {
    "深蓝": "1F4E79",
    "蓝绿": "008C95",
    "蓝": "1F4E79",
    "绿色": "00A65A",
    "绿": "00A65A",
    "青": "0099A8",
    "灰": "D9E2F3",
    "红": "C00000",
    "黑": "000000",
}

FONT_KEYWORDS = {
    "仿宋": ("FangSong", "仿宋"),
    "宋体": ("SimSun", "宋体"),
    "黑体": ("SimHei", "黑体"),
    "楷体": ("KaiTi", "楷体"),
    "微软雅黑": ("Microsoft YaHei", "微软雅黑"),
    "serif": ("Times New Roman", "Times New Roman"),
    "times": ("Times New Roman", "Times New Roman"),
}


def safe_filename(value: str, fallback: str = "draft") -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\r\n\t]+', "_", value).strip(" ._")
    return cleaned or fallback


def build_style_config(layout_profile: dict[str, Any] | None, *, title: str) -> dict[str, Any]:
    profile = layout_profile or {}
    body_font = _infer_font(
        _profile_values(profile, "body", "正文", "normal", "fonts", "style_rules"),
        default=("FangSong", "仿宋"),
    )
    heading_font = _infer_font(
        _profile_values(profile, "heading", "标题", "heading_hierarchy", "fonts", "style_rules"),
        default=("SimHei", "黑体"),
    )
    cover_font = _infer_font(
        _profile_values(profile, "cover", "封面", "fonts", "style_rules"),
        default=heading_font,
    )
    table_font = _infer_font(
        _profile_values(profile, "table", "表格", "tables", "fonts", "style_rules"),
        default=("SimSun", "宋体"),
    )
    heading_color = _infer_color(
        _heading_color_values(profile),
        default="1F4E79",
    )
    cover_color = _infer_color(
        _profile_values(profile, "cover_color", "cover", "封面", "style_rules"),
        default=heading_color,
    )
    cover_accent_fill = _infer_color(
        _profile_values(profile, "accent", "wave", "波浪", "cover", "封面", "style_rules"),
        default="EAF4F8",
    )
    table_header_fill = _infer_color(
        _profile_values(profile, "header_fill", "table_header", "表头", "tables", "style_rules"),
        default="D9EAF7",
    )
    table_header_text_color = _infer_color(
        _profile_values(profile, "table_header_text", "表头文字", "tables", "style_rules"),
        default="000000",
    )
    return {
        "title": title,
        "body_font": body_font,
        "heading_font": heading_font,
        "cover_font": cover_font,
        "table_font": table_font,
        "heading_color": heading_color,
        "cover_color": cover_color,
        "cover_accent_fill": cover_accent_fill,
        "table_header_fill": table_header_fill,
        "table_header_text_color": table_header_text_color,
        "header_text": _profile_header_text(profile, title=title),
        "cover_subtitle": _layout_text(profile, "cover", "subtitle") or _find_text(profile, ["subtitle", "副标题"]) or "通过制技术标",
        "cover_bottom_text": _layout_text(profile, "cover", "bottom_text") or _find_text(profile, ["bottom_text", "底部"]) or "投标文件技术部分",
        "numbering_scheme": _infer_numbering_scheme(profile),
    }


def apply_chinese_font(document: Document, style_config: dict[str, Any]) -> None:
    styles = document.styles
    body_font, body_east_asia = style_config["body_font"]
    heading_font, heading_east_asia = style_config["heading_font"]
    normal = styles["Normal"]
    normal.font.name = body_font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_east_asia)
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.5
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = heading_font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), heading_east_asia)
        style.font.bold = True
        style.font.color.rgb = _rgb(style_config["heading_color"])
        style.paragraph_format.first_line_indent = Pt(0)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(14)


def apply_layout_profile(
    document: Document,
    *,
    style_config: dict[str, Any],
) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.45)

    header_text = style_config["header_text"]
    paragraph = section.header.paragraphs[0]
    paragraph.text = header_text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_font(paragraph, "SimSun", 9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    _add_page_number(footer)
    footer.add_run(" 页")
    _set_paragraph_font(footer, "SimSun", 9)


def add_cover_page(document: Document, title: str, style_config: dict[str, Any]) -> None:
    add_cover_accent_band(document, style_config)
    for _ in range(5):
        document.add_paragraph("")
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.first_line_indent = Pt(0)
    title_run = title_para.add_run(title)
    title_run.bold = True
    cover_font, cover_east_asia = style_config["cover_font"]
    _set_run_font(title_run, cover_font, cover_east_asia, 22, color=style_config["cover_color"])

    subtitle = style_config["cover_subtitle"]
    subtitle_para = document.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_para.paragraph_format.first_line_indent = Pt(0)
    subtitle_run = subtitle_para.add_run(subtitle)
    _set_run_font(subtitle_run, cover_font, cover_east_asia, 16, color=style_config["cover_color"])

    for _ in range(8):
        document.add_paragraph("")
    note_para = document.add_paragraph()
    note_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_para.paragraph_format.first_line_indent = Pt(0)
    note = style_config["cover_bottom_text"]
    note_run = note_para.add_run(note)
    _set_run_font(note_run, "SimSun", "宋体", 12)
    document.add_page_break()


def add_cover_accent_band(document: Document, style_config: dict[str, Any]) -> None:
    paragraph = document.add_paragraph(" ")
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), style_config["cover_accent_fill"])
    p_pr.append(shading)


def add_toc_page(document: Document) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Pt(0)
    run = heading.add_run("目录")
    run.bold = True
    _set_run_font(run, "SimHei", "黑体", 16)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _add_complex_field(paragraph, r'TOC \o "1-3" \h \z \u')
    document.add_page_break()


def _profile_header_text(profile: dict[str, Any], *, title: str) -> str:
    explicit = _layout_text_or_none(profile, "headers_footers", "header")
    if explicit is None:
        explicit = _find_text(profile, ["页眉", "header"]) or None
    if explicit is None:
        return title
    return "" if _is_absent_layout_text(explicit) else explicit


def _layout_text(profile: dict[str, Any], section: str, key: str) -> str:
    return _layout_text_or_none(profile, section, key) or ""


def _layout_text_or_none(profile: dict[str, Any], section: str, key: str) -> str | None:
    value = profile.get(section, {})
    if isinstance(value, dict):
        candidate = value.get(key) or value.get("text") or value.get("pattern")
        if isinstance(candidate, str):
            return candidate.strip()
    return None


def _is_absent_layout_text(value: str) -> bool:
    normalized = re.sub(r"[\s:：。；;，,、（）()]+", "", value).lower()
    return normalized in {
        "",
        "无",
        "没有",
        "无页眉",
        "无页脚",
        "不设置",
        "未设置",
        "none",
        "null",
        "no",
        "n/a",
        "na",
    }


def _heading_color_values(profile: dict[str, Any]) -> list[str]:
    values = _profile_values_by_path(
        profile,
        "heading_color",
        "title_color",
        "标题颜色",
        "heading_hierarchy",
        "fonts.heading",
        "fonts.heading_font",
        "fonts.title",
        "fonts.title_font",
    )
    values.extend(_style_rule_values(profile, "标题", "heading", "title"))
    return values


def _profile_values_by_path(profile: Any, *hints: str) -> list[str]:
    hints_lower = [hint.lower() for hint in hints]
    values: list[str] = []

    def collect(value: Any) -> None:
        if isinstance(value, dict):
            for child in value.values():
                collect(child)
        elif isinstance(value, list):
            for item in value:
                collect(item)
        else:
            text = _stringify(value).strip()
            if text:
                values.append(text)

    def walk(value: Any, path: str = "") -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                key_text = str(key)
                child_path = f"{path}.{key_text}" if path else key_text
                if any(hint in child_path.lower() for hint in hints_lower):
                    collect(child)
                walk(child, child_path)
        elif isinstance(value, list):
            for item in value:
                walk(item, path)

    walk(profile)
    return values


def _style_rule_values(profile: dict[str, Any], *text_hints: str) -> list[str]:
    rules = profile.get("style_rules", [])
    if isinstance(rules, str):
        rules = [rules]
    if not isinstance(rules, list):
        return []
    hints_lower = [hint.lower() for hint in text_hints]
    values: list[str] = []
    for rule in rules:
        text = _stringify(rule).strip()
        if not text:
            continue
        clauses = [clause.strip() for clause in re.split(r"[，,；;。]", text) if clause.strip()]
        matches = [clause for clause in clauses if any(hint in clause.lower() for hint in hints_lower)]
        values.extend(matches or ([text] if any(hint in text.lower() for hint in hints_lower) else []))
    return values


def _infer_numbering_scheme(profile: dict[str, Any]) -> dict[str, Any]:
    values = _profile_values_by_path(
        profile,
        "numbering",
        "编号",
        "编号体系",
        "heading_hierarchy",
        "章节编号",
    )
    text = " ".join(values)
    chinese_heading = bool(
        re.search(r"第[一二三四五六七八九十百\d]+章|第一章|第X章|第x章", text)
        or re.search(r"第[一二三四五六七八九十百\d]+节|第一节|第X节|第x节", text)
        or "一、" in text
    )
    if not chinese_heading:
        return {"enabled": False}
    return {
        "enabled": True,
        "level_1": "chapter",
        "level_2": "section",
        "level_3": "chinese_comma",
        "level_4": "decimal_comma",
    }


def _profile_values(profile: Any, *hints: str) -> list[str]:
    hints_lower = [hint.lower() for hint in hints]
    values: list[str] = []

    def walk(value: Any, path: str = "") -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                key_text = str(key)
                child_path = f"{path}.{key_text}" if path else key_text
                if any(hint in key_text.lower() or hint in child_path.lower() for hint in hints_lower):
                    values.append(_stringify(child))
                walk(child, child_path)
        elif isinstance(value, list):
            for item in value:
                walk(item, path)
        elif isinstance(value, str):
            if any(hint in value.lower() or hint in path.lower() for hint in hints_lower):
                values.append(value)

    walk(profile)
    return [value for value in values if value.strip()]


def _find_text(profile: Any, hints: list[str]) -> str:
    values = _profile_values(profile, *hints)
    for value in values:
        stripped = re.sub(r"\s+", " ", value).strip()
        if 1 <= len(stripped) <= 80 and not stripped.startswith("{"):
            return stripped
    return ""


def _stringify(value: Any) -> str:
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    return str(value)


def _infer_font(values: list[str], *, default: tuple[str, str]) -> tuple[str, str]:
    text = " ".join(values).lower()
    for keyword, font in FONT_KEYWORDS.items():
        if keyword.lower() in text:
            return font
    return default


def _infer_color(values: list[str], *, default: str) -> str:
    text = " ".join(values)
    explicit = re.search(r"#?([0-9A-Fa-f]{6})", text)
    if explicit:
        return explicit.group(1).upper()
    rgb = re.search(r"rgb\((\d{1,3}),\s*(\d{1,3}),\s*(\d{1,3})\)", text, re.I)
    if rgb:
        return "".join(f"{max(0, min(255, int(part))):02X}" for part in rgb.groups())
    for keyword, color in COLOR_KEYWORDS.items():
        if keyword in text:
            return color
    return default


def _rgb(color: str) -> RGBColor:
    clean = re.sub(r"[^0-9A-Fa-f]", "", color or "")
    if len(clean) != 6:
        clean = "000000"
    return RGBColor(int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16))


def _set_paragraph_font(paragraph, font_name: str, size: int) -> None:
    for run in paragraph.runs:
        _set_run_font(run, font_name, font_name, size)


def _set_run_font(run, font_name: str, east_asia: str, size: int | float, color: str | None = None) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = _rgb(color)


def _add_page_number(paragraph) -> None:
    _add_complex_field(paragraph, "PAGE")


def _add_complex_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    style_config: dict[str, Any] | None = None,
) -> None:
    style_config = style_config or {}
    table_font, table_east_asia = style_config.get("table_font", ("SimSun", "宋体"))
    text_color = style_config.get("table_header_text_color", "000000") if bold else None
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        segments = re.split(r"(\*\*.+?\*\*)", str(text))
        for segment in segments:
            if not segment:
                continue
            marked_bold = segment.startswith("**") and segment.endswith("**")
            clean = segment[2:-2] if marked_bold else segment
            run = paragraph.add_run(clean)
            run.bold = bold or marked_bold
            _set_run_font(run, table_font, table_east_asia, 10.5, color=text_color)


def _set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _apply_table_geometry(table, rows: list[list[str]], total_width: int = 8640) -> None:
    if not rows or not table.columns:
        return
    column_count = len(table.columns)
    weights: list[int] = []
    for column in range(column_count):
        lengths = [
            len(str(row[column])) if column < len(row) else 0
            for row in rows
        ]
        weights.append(max(6, min(max(lengths, default=6), 40)))
    weight_total = max(sum(weights), 1)
    minimum_width = max(360, min(720, total_width // column_count))
    widths = [max(minimum_width, round(total_width * weight / weight_total)) for weight in weights]
    difference = total_width - sum(widths)
    widths[-1] += difference
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.first_child_found_in("w:tblW")
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(total_width))
    tbl_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for index, grid_col in enumerate(grid.gridCol_lst):
        grid_col.set(qn("w:w"), str(widths[index]))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def normalize_sections(
    *,
    outline: list[dict[str, Any]] | None,
    sections: dict[str, str] | list[dict[str, Any]] | None,
    response_matrix: list[dict[str, Any]] | None,
) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    if isinstance(sections, dict):
        for idx, (title, content) in enumerate(sections.items(), start=1):
            normalized.append({"level": 1, "title": title, "content": content, "order": idx})
    elif isinstance(sections, list):
        for idx, item in enumerate(sections, start=1):
            normalized.append(
                {
                    "level": int(item.get("level", 1)),
                    "title": str(item.get("title", f"章节{idx}")),
                    "content": str(item.get("content", "")),
                    "order": int(item.get("order", idx)),
                }
            )
    elif outline:
        matrix_by_section: dict[str, list[dict[str, Any]]] = {}
        for row in response_matrix or []:
            matrix_by_section.setdefault(str(row.get("target_section", "")), []).append(row)
        for idx, item in enumerate(outline, start=1):
            title = str(item.get("title", f"章节{idx}"))
            rows = matrix_by_section.get(title, [])
            if rows:
                body = "\n".join(
                    f"【待完善】响应 {row.get('requirement_id')}: {row.get('response_strategy')}"
                    for row in rows
                )
            else:
                body = "【待完善】请 Hermes 根据响应矩阵和历史通过案例补写本章内容。"
            normalized.append(
                {
                    "level": int(item.get("level", 1)),
                    "title": title,
                    "content": body,
                    "order": int(item.get("order", idx)),
                }
            )
    else:
        normalized.append(
            {
                "level": 1,
                "title": "施工组织设计",
                "content": "【待完善】请 Hermes 根据招标文件要求生成通过制技术标正文。",
                "order": 1,
            }
        )
    return normalized


def add_markdown_content(
    document: Document,
    content: str,
    *,
    style_config: dict[str, Any],
    numbering_state: dict[str, Any] | None = None,
) -> dict[str, int]:
    lines = content.splitlines()
    idx = 0
    numbered_list_counter = 0
    stats = {"paragraphs": 0, "tables": 0, "headings": 0, "lists": 0}
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            numbered_list_counter = 0
            idx += 1
            continue
        if _is_markdown_table_start(lines, idx):
            rows, idx = _consume_markdown_table(lines, idx)
            add_word_table(document, rows, style_config=style_config)
            numbered_list_counter = 0
            stats["tables"] += 1
            continue
        image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if image:
            image_path = Path(image.group(2).strip().strip('"')).expanduser()
            if image_path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
                if image.group(1).strip():
                    caption = document.add_paragraph(image.group(1).strip())
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.runs[0].italic = True
                stats["paragraphs"] += 1
            else:
                document.add_paragraph(f"【待人工确认：图形文件不存在：{image_path}】")
            idx += 1
            continue
        heading = _parse_markdown_or_chinese_heading(line)
        if heading:
            level, title = heading
            add_profile_heading(
                document,
                title,
                level=level,
                style_config=style_config,
                numbering_state=numbering_state,
            )
            numbered_list_counter = 0
            stats["headings"] += 1
            idx += 1
            continue
        if re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_markdown(paragraph, re.sub(r"^[-*]\s+", "", line))
            paragraph.paragraph_format.first_line_indent = Pt(0)
            numbered_list_counter = 0
            stats["lists"] += 1
            idx += 1
            continue
        numbered = re.match(r"^\d+[.)、]\s+(.+)$", line)
        if numbered:
            numbered_list_counter += 1
            paragraph = document.add_paragraph()
            _add_inline_markdown(
                paragraph,
                f"{numbered_list_counter}、{numbered.group(1).strip()}",
            )
            paragraph.paragraph_format.first_line_indent = Pt(0)
            stats["lists"] += 1
            idx += 1
            continue
        paragraph = document.add_paragraph()
        _add_inline_markdown(paragraph, line)
        numbered_list_counter = 0
        stats["paragraphs"] += 1
        idx += 1
    return stats


def _add_inline_markdown(paragraph, text: str) -> None:
    for segment in re.split(r"(\*\*.+?\*\*)", text):
        if not segment:
            continue
        marked_bold = segment.startswith("**") and segment.endswith("**")
        clean = segment[2:-2] if marked_bold else segment
        run = paragraph.add_run(clean)
        run.bold = marked_bold


def _parse_markdown_or_chinese_heading(line: str) -> tuple[int, str] | None:
    markdown = re.match(r"^(#{1,4})\s+(.+?)\s*$", line)
    if markdown:
        return len(markdown.group(1)), markdown.group(2).strip()
    if re.match(r"^第[一二三四五六七八九十百\d]+章[、\s：:]", line):
        return 1, line
    if re.match(r"^第[一二三四五六七八九十百\d]+节[、\s：:]", line):
        return 2, line
    if re.match(r"^[一二三四五六七八九十]+[、.．]\s*", line):
        return 3, line
    return None


def _new_numbering_state(style_config: dict[str, Any]) -> dict[str, Any]:
    scheme = style_config.get("numbering_scheme", {})
    return {
        "enabled": bool(isinstance(scheme, dict) and scheme.get("enabled")),
        "counters": [0, 0, 0, 0],
    }


def add_profile_heading(
    document: Document,
    title: str,
    *,
    level: int,
    style_config: dict[str, Any],
    numbering_state: dict[str, Any] | None,
) -> None:
    display_title = _format_profile_heading(title, level=level, numbering_state=numbering_state)
    document.add_heading(display_title, level=max(1, min(level, 3)))


def _format_profile_heading(
    title: str,
    *,
    level: int,
    numbering_state: dict[str, Any] | None,
) -> str:
    normalized_level = max(1, min(level, 4))
    if not numbering_state or not numbering_state.get("enabled"):
        return title

    counters = numbering_state["counters"]
    counters[normalized_level - 1] += 1
    for idx in range(normalized_level, len(counters)):
        counters[idx] = 0

    bare_title = _strip_existing_heading_number(title)
    if normalized_level == 1:
        return f"第{_to_chinese_number(counters[0])}章 {bare_title}"
    if normalized_level == 2:
        return f"第{_to_chinese_number(counters[1])}节 {bare_title}"
    if normalized_level == 3:
        return f"{_to_chinese_number(counters[2])}、{bare_title}"
    return f"{counters[3]}、{bare_title}"


def _strip_existing_heading_number(title: str) -> str:
    text = title.strip()
    patterns = [
        r"^第[一二三四五六七八九十百\d]+[章节][、\s：:.-]*",
        r"^[一二三四五六七八九十百]+[、.．]\s*",
        r"^\d+(?:\.\d+)*(?:[.)、．]|\s+)",
    ]
    for pattern in patterns:
        stripped = re.sub(pattern, "", text, count=1).strip()
        if stripped != text and stripped:
            return stripped
    return text


def _to_chinese_number(value: int) -> str:
    numerals = "零一二三四五六七八九"
    if value <= 0:
        return str(value)
    if value < 10:
        return numerals[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + numerals[value % 10]
    if value < 100:
        tens, ones = divmod(value, 10)
        return numerals[tens] + "十" + (numerals[ones] if ones else "")
    return str(value)


def _is_markdown_table_start(lines: list[str], idx: int) -> bool:
    if idx + 1 >= len(lines):
        return False
    return _is_pipe_row(lines[idx]) and _is_table_separator(lines[idx + 1])


def _is_pipe_row(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not _is_pipe_row(stripped):
        return False
    body = stripped.strip("|").strip()
    return bool(body) and all(re.fullmatch(r":?-{3,}:?", part.strip()) for part in body.split("|"))


def _consume_markdown_table(lines: list[str], idx: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while idx < len(lines) and _is_pipe_row(lines[idx]):
        if not _is_table_separator(lines[idx]):
            rows.append([cell.strip() for cell in lines[idx].strip().strip("|").split("|")])
        idx += 1
    width = max((len(row) for row in rows), default=0)
    normalized = [row + [""] * (width - len(row)) for row in rows if width]
    return normalized, idx


def add_word_table(
    document: Document,
    rows: list[list[str]],
    *,
    style_config: dict[str, Any],
) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_fill = style_config["table_header_fill"]
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for col_idx in range(column_count):
            cell = cells[col_idx]
            text = row[col_idx] if col_idx < len(row) else ""
            _set_cell_text(cell, text, bold=row_idx == 0, style_config=style_config)
            if row_idx == 0:
                _shade_cell(cell, header_fill)
    _apply_table_geometry(table, rows)


def generate_docx(
    *,
    title: str,
    output_path: Path,
    outline: list[dict[str, Any]] | None = None,
    sections: dict[str, str] | list[dict[str, Any]] | None = None,
    requirements: dict[str, Any] | None = None,
    response_matrix: list[dict[str, Any]] | None = None,
    layout_profile: dict[str, Any] | None = None,
    include_requirement_response: bool = False,
) -> dict[str, Any]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    style_config = build_style_config(layout_profile, title=title)
    apply_chinese_font(document, style_config)
    apply_layout_profile(document, style_config=style_config)
    add_cover_page(document, title, style_config)
    add_toc_page(document)
    document.add_paragraph("说明：本文件为 Hermes 生成的通过制技术标复核初稿，提交前需人工核查项目参数、格式和签章要求。")
    numbering_state = _new_numbering_state(style_config)

    if requirements:
        document.add_heading("通过制响应摘要", level=1)
        document.add_paragraph(f"检测模式：{requirements.get('detected_mode', 'unknown')}")
        document.add_paragraph(f"抽取要求数量：{requirements.get('requirement_count', 0)}")
        if include_requirement_response:
            document.add_heading("招标要求逐条响应说明", level=2)
            for item in requirements.get("requirements", []):
                requirement = str(item.get("requirement", "")).strip()
                if not requirement:
                    continue
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{item.get('id', '')}：").bold = True
                paragraph.add_run(requirement)
                paragraph.add_run(" 本文件已在对应章节响应，最终以响应矩阵和人工复核结果为准。")

    for section in normalize_sections(
        outline=outline,
        sections=sections,
        response_matrix=response_matrix,
    ):
        level = max(1, min(int(section.get("level", 1)), 3))
        add_profile_heading(
            document,
            str(section["title"]),
            level=level,
            style_config=style_config,
            numbering_state=numbering_state,
        )
        content = str(section.get("content", "")).strip()
        if content:
            add_markdown_content(
                document,
                content,
                style_config=style_config,
                numbering_state=numbering_state,
            )

    if response_matrix:
        document.add_heading("响应矩阵复核表", level=1)
        table = document.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        headers = ["编号", "类别", "招标要求", "对应章节", "人工核查"]
        for idx, header in enumerate(headers):
            _set_cell_text(table.rows[0].cells[idx], header, bold=True, style_config=style_config)
            _shade_cell(table.rows[0].cells[idx], style_config["table_header_fill"])
        for row in response_matrix:
            cells = table.add_row().cells
            _set_cell_text(cells[0], str(row.get("requirement_id", "")), style_config=style_config)
            _set_cell_text(cells[1], str(row.get("category", "")), style_config=style_config)
            _set_cell_text(cells[2], str(row.get("requirement", "")), style_config=style_config)
            _set_cell_text(cells[3], str(row.get("target_section", "")), style_config=style_config)
            _set_cell_text(cells[4], "是" if row.get("human_check") else "否", style_config=style_config)
        _apply_table_geometry(
            table,
            [
                headers,
                *[
                    [
                        str(row.get("requirement_id", "")),
                        str(row.get("category", "")),
                        str(row.get("requirement", "")),
                        str(row.get("target_section", "")),
                        "是" if row.get("human_check") else "否",
                    ]
                    for row in response_matrix
                ],
            ],
        )

    document.save(str(output_path))
    field_update = update_docx_fields(output_path)
    return {
        "docx_path": str(output_path),
        "section_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "field_update": field_update,
        "style_config": style_config,
    }


def update_docx_fields(docx_path: Path) -> dict[str, Any]:
    """Update TOC/page fields when Microsoft Word COM is available."""
    if os.getenv("PYTEST_CURRENT_TEST") or os.getenv("PASS_BID_WRITING_DISABLE_WORD_COM") == "1":
        return {
            "status": "skipped",
            "message": "Word COM field update is disabled for deterministic test execution.",
        }
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        return {"status": "skipped", "engine": "none", "message": str(exc)}

    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        document = word.Documents.Open(str(docx_path.resolve()))
        document.Fields.Update()
        for toc in document.TablesOfContents:
            toc.Update()
        document.Save()
        document.Close(False)
        return {"status": "ready", "engine": "word_com"}
    except Exception as exc:
        return {"status": "blocked", "engine": "word_com", "message": str(exc)}
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def check_docx_compliance(
    *,
    docx_path: Path,
    requirements: dict[str, Any],
) -> dict[str, Any]:
    draft_text = _extract_docx_paragraph_text(docx_path)
    missing: list[dict[str, Any]] = []
    covered: list[dict[str, Any]] = []
    for item in requirements.get("requirements", []):
        requirement = str(item.get("requirement", ""))
        keywords = [kw for kw in re.split(r"[，。；、\s]+", requirement) if len(kw) >= 2]
        important = keywords[:8]
        hits = [kw for kw in important if kw in draft_text]
        result = {
            "requirement_id": item.get("id"),
            "category": item.get("category"),
            "requirement": requirement,
            "matched_keywords": hits,
        }
        if hits:
            covered.append(result)
        else:
            missing.append(result)

    placeholders = sorted(set(re.findall(r"【[^】]*(?:待完善|人工确认|待确认)[^】]*】", draft_text)))
    return {
        "docx_path": str(docx_path),
        "checked_text_scope": "paragraphs_only",
        "requirement_count": len(requirements.get("requirements", [])),
        "covered_count": len(covered),
        "missing_count": len(missing),
        "placeholder_count": len(placeholders),
        "missing_requirements": missing,
        "placeholders": placeholders,
        "status": "needs_human_review" if missing or placeholders else "ready_for_manual_final_check",
    }


def _extract_docx_paragraph_text(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    parts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def export_docx_to_pdf(docx_path: Path, pdf_path: Path) -> dict[str, Any]:
    docx_path = docx_path.resolve()
    pdf_path = pdf_path.resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    libreoffice = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice:
        completed = subprocess.run(
            [
                libreoffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_path.parent),
                str(docx_path),
            ],
            text=True,
            capture_output=True,
            timeout=120,
            check=False,
        )
        generated = pdf_path.parent / f"{docx_path.stem}.pdf"
        if generated.exists() and generated != pdf_path:
            generated.replace(pdf_path)
        if pdf_path.exists():
            return {"status": "ready", "pdf_path": str(pdf_path), "engine": "libreoffice"}
        return {
            "status": "blocked",
            "engine": "libreoffice",
            "stdout": completed.stdout,
            "stderr": completed.stderr,
            "message": "LibreOffice conversion did not create the expected PDF.",
        }

    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        return {
            "status": "blocked",
            "pdf_path": str(pdf_path),
            "engine": "none",
            "message": f"PDF export requires Microsoft Word COM or LibreOffice: {exc}",
        }

    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        document = word.Documents.Open(str(docx_path))
        document.Fields.Update()
        for toc in document.TablesOfContents:
            toc.Update()
        document.SaveAs(str(pdf_path), FileFormat=17)
        document.Close(False)
        return {"status": "ready", "pdf_path": str(pdf_path), "engine": "word_com"}
    except Exception as exc:
        return {
            "status": "blocked",
            "pdf_path": str(pdf_path),
            "engine": "word_com",
            "message": str(exc),
        }
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
