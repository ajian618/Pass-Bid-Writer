from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from docx import Document
from PIL import Image

from pass_bid_writing.assets import generate_document_assets
from pass_bid_writing.documents import generate_docx
from pass_bid_writing.production import (
    _sanitize_generated_structure,
    _validate_review_issues,
    confirm_file_roles,
    confirm_task_spec,
    extract_standards,
    get_run_state,
    prepare_production_project,
    _sanitize_unsupported_claims,
)
from pass_bid_writing.visual_sources import _render_dxf_preview, analyze_visual_sources
from pass_bid_writing.knowledge import attach_standard_file
from pass_bid_writing import db
from pass_bid_writing.analysis import extract_case_patterns


class ProductionSystemTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.storage = self.root / "storage"
        self.project = self.root / "projects" / "new_tenders" / "河道治理测试项目"
        self.project.mkdir(parents=True)
        self.previous_storage = os.environ.get("PASS_BID_WRITING_STORAGE_DIR")
        os.environ["PASS_BID_WRITING_STORAGE_DIR"] = str(self.storage)

    def tearDown(self) -> None:
        if self.previous_storage is None:
            os.environ.pop("PASS_BID_WRITING_STORAGE_DIR", None)
        else:
            os.environ["PASS_BID_WRITING_STORAGE_DIR"] = self.previous_storage
        self.tempdir.cleanup()

    def test_reference_and_evidence_production_flow(self) -> None:
        (self.project / "招标文件.txt").write_text(
            "\n".join(
                [
                    "某河道治理项目施工组织设计采用通过制评审。",
                    "计划总工期：180日历天。",
                    "质量目标：合格。",
                    "技术标准和要求：《水利水电工程施工组织设计规范》SL 303-2017。",
                    "施工质量验收执行《水利水电建设工程验收规程》SL 223-2008。",
                    "必须提供施工进度计划、主要施工方案、质量安全措施、人员机械配置和总平面布置。",
                ]
            ),
            encoding="utf-8",
        )
        (self.project / "初步设计报告.txt").write_text(
            "项目名称：某河道治理项目\n建设地点：浙江省某市\n质量目标：合格\n",
            encoding="utf-8",
        )
        (self.project / "招标预算.txt").write_text("工程量清单：河道疏浚、护岸工程。", encoding="utf-8")
        (self.project / "设计图纸.dxf").write_text("0\nSECTION\n2\nENTITIES\n0\nENDSEC\n0\nEOF", encoding="ascii")

        state = prepare_production_project(str(self.project), project_type="河道治理工程")

        self.assertIsInstance(state["run_id"], int)
        self.assertGreaterEqual(len(state["standards"]), 2)
        self.assertTrue(all(item["status"] == "pending_download" for item in state["standards"]))
        self.assertGreaterEqual(len(state["facts"]), 3)
        self.assertEqual(state["blueprint"]["title"], "水利工程施工组织设计标准蓝图")
        self.assertEqual(len(state["sections"]), 15)
        self.assertTrue(any("河道疏浚" in item for item in state["sections"][5]["components"]))
        self.assertTrue(any(item["category"] == "规范待下载" for item in state["confirmations"]))

        expected_reports = {
            "task_spec",
            "response_matrix",
            "basis_report",
            "standards_queue",
            "confirmations",
            "model_differences",
        }
        self.assertEqual(set(state["reports"]), expected_reports)
        for report in state["reports"].values():
            self.assertTrue(Path(report).exists())
            workbook = load_workbook(report, read_only=True)
            self.assertGreaterEqual(len(workbook.sheetnames), 1)
            workbook.close()

        roles = confirm_file_roles(state["run_id"])
        self.assertEqual(roles["status"], "visual_analysis_pending")
        visual = analyze_visual_sources(state["run_id"])
        self.assertEqual(visual["status"], "awaiting_confirmation")
        self.assertTrue(any(item["category"] == "视觉模型未配置" for item in visual["confirmations"]))

        confirmed = confirm_task_spec(state["run_id"])
        self.assertEqual(confirmed["status"], "ready_for_generation")
        reloaded = get_run_state(state["run_id"])
        self.assertEqual(reloaded["stage"], "章节生成")

    def test_standard_extraction_keeps_source_page_and_official_queue(self) -> None:
        standards = extract_standards(
            [
                {
                    "page": 28,
                    "text": "本工程执行《水利水电工程施工组织设计规范》SL 303-2017。",
                }
            ],
            "招标文件.pdf",
        )
        self.assertEqual(standards[0]["code"], "SL 303-2017")
        self.assertEqual(standards[0]["source_page"], 28)
        self.assertIn("水利部", standards[0]["official_platform"])
        self.assertEqual(standards[0]["status"], "pending_download")

    def test_generated_assets_are_editable_and_embeddable(self) -> None:
        state = {
            "project": {"name": "测试项目"},
            "facts": [],
        }
        assets = generate_document_assets(state, self.root / "outputs")
        self.assertTrue(Path(assets["production_workflow_svg"]).exists())
        self.assertTrue(Path(assets["schedule_source_json"]).exists())
        image_path = Path(assets["quality_flow_png"])
        with Image.open(image_path) as image:
            self.assertGreater(image.width, 1000)

        docx_path = self.root / "outputs" / "图形嵌入测试.docx"
        generate_docx(
            title="图形嵌入测试",
            output_path=docx_path,
            sections=[
                {
                    "level": 1,
                    "title": "质量控制",
                    "content": f"质量控制正文。\n\n![质量控制流程图]({image_path})",
                }
            ],
        )
        document = Document(docx_path)
        self.assertEqual(len(document.inline_shapes), 1)

    def test_unsupported_numbers_and_clause_citations_are_blocked(self) -> None:
        state = {
            "facts": [{"key": "总工期", "value": "180", "unit": "日历天"}],
            "standard_clauses": [],
        }
        section = {
            "requirements": [],
            "project_basis": ["总工期：180日历天"],
        }
        content, findings = _sanitize_unsupported_claims(
            "总工期180日历天；压实度95%；依据SL 260-2014第6.2条。",
            state,
            section,
        )
        self.assertIn("180日历天", content)
        self.assertIn("【待人工确认：95%的项目或规范依据】", content)
        self.assertIn("【待人工确认：SL 260-2014第6.2条的项目或规范依据】", content)
        self.assertEqual(len(findings), 2)

    def test_uploaded_standard_is_parsed_into_verified_clauses(self) -> None:
        (self.project / "招标文件.txt").write_text(
            "技术标准和要求：《水利水电工程施工组织设计规范》SL 303-2017。",
            encoding="utf-8",
        )
        state = prepare_production_project(str(self.project))
        standard_file = self.project / "SL303规范原文.txt"
        standard_file.write_text(
            "1.0.1 本规范适用于水利水电工程施工组织设计。\n"
            "3.2.1 施工总布置应结合施工条件统筹确定。\n"
            "3.2.2 严禁无依据改变设计边界。\n",
            encoding="utf-8",
        )
        result = attach_standard_file(
            run_id=state["run_id"],
            state=state,
            standard_code="SL 303-2017",
            target=standard_file,
        )
        self.assertEqual(len(result["clauses"]), 3)
        self.assertEqual(result["clauses"][2]["mandatory_level"], "mandatory")
        with db.db_session(self.storage / "pass_bid_writing.db") as conn:
            count = conn.execute(
                "SELECT COUNT(*) FROM standard_clauses WHERE production_run_id = ?",
                (state["run_id"],),
            ).fetchone()[0]
        self.assertEqual(count, 3)

    def test_generation_persists_content_and_runs_mandatory_audits(self) -> None:
        (self.project / "招标文件.txt").write_text(
            "某河道工程施工组织设计采用通过制。总工期180日历天。"
            "必须提供进度、质量、安全、机械、人员和主要施工方案。",
            encoding="utf-8",
        )
        state = prepare_production_project(str(self.project), project_type="河道治理工程")
        confirm_file_roles(state["run_id"])
        confirmed = confirm_task_spec(state["run_id"])
        self.assertEqual(confirmed["status"], "ready_for_generation")

        def fake_complete_json(*, prompt: str, **_: object) -> dict[str, object]:
            if "最终文字复核员" in str(_.get("system", "")):
                return {"status": "passed", "issues": []}
            return {
                "content": "## 编制依据\n本章仅依据项目证据编制。\n\n| 控制项 | 要求 |\n|---|---|\n| 项目参数 | 【待人工确认】 |",
                "human_confirm_items": [],
            }

        with (
            patch("pass_bid_writing.production.ModelRouter.complete_json", side_effect=fake_complete_json),
            patch(
                "pass_bid_writing.production.visual_check_document",
                return_value={"status": "disabled", "message": "test"},
            ),
        ):
            generated = __import__(
                "pass_bid_writing.production",
                fromlist=["generate_production_docx"],
            ).generate_production_docx(state["run_id"])
        self.assertTrue(all(section.get("content") for section in generated["sections"]))
        self.assertIn("compliance", generated)
        self.assertIn("document_audit", generated)
        self.assertGreaterEqual(generated["document_audit"]["image_count"], 4)
        reloaded = get_run_state(state["run_id"])
        self.assertTrue(all(section.get("content") for section in reloaded["sections"]))

    def test_accepted_case_assets_enter_reference_basis_without_parameters(self) -> None:
        settings_db = self.storage / "pass_bid_writing.db"
        db.init_db(settings_db)
        tender_text = "河道治理工程，要求施工进度、质量和安全措施。"
        bid_text = "第一章 施工总体部署\n围堰长25m，导流施工按设计图纸组织，具体水位以本项目设计为准。"
        patterns = extract_case_patterns(tender_text, bid_text)
        with db.db_session(settings_db) as conn:
            db.create_case_pair(
                conn,
                title="已通过河道案例",
                project_type="河道治理",
                region="浙江",
                tags="河道,通过制",
                tender_path="历史招标文件.pdf",
                bid_path="历史已通过技术标.docx",
                tender_text=tender_text,
                bid_text=bid_text,
                outline=patterns["outline"],
                requirements={},
                patterns=patterns,
            )
        (self.project / "招标文件.txt").write_text(
            "河道治理工程施工组织设计采用通过制，要求围堰导流、施工进度和质量措施。",
            encoding="utf-8",
        )
        state = prepare_production_project(str(self.project), project_type="河道治理工程")
        self.assertGreater(len(state["case_assets"]), 0)
        self.assertTrue(
            any(
                "案例资产：" in basis
                for section in state["sections"]
                for basis in section["reference_basis"]
            )
        )
        self.assertFalse(any(item["category"] == "案例资产缺失" for item in state["confirmations"]))
        wording = [asset for asset in state["case_assets"] if asset["asset_type"] == "wording"]
        self.assertTrue(all("25m" not in str(asset["content"]) for asset in wording))

    def test_qwen_visual_results_become_located_project_facts(self) -> None:
        (self.project / "招标文件.txt").write_text(
            "河道工程施工组织设计采用通过制。",
            encoding="utf-8",
        )
        Image.new("RGB", (800, 600), "white").save(self.project / "设计图纸.png")
        state = prepare_production_project(str(self.project), project_type="河道治理工程")
        roles = confirm_file_roles(state["run_id"])
        self.assertEqual(roles["status"], "visual_analysis_pending")
        statuses = [
            {"role": "text_master", "provider": "deepseek", "model": "deepseek-v4-pro", "configured": True, "modalities": ["text"]},
            {"role": "vision_primary", "provider": "aliyun_qwen", "model": "qwen-test", "configured": True, "modalities": ["text", "image"]},
            {"role": "vision_batch", "provider": "aliyun_qwen", "model": "qwen-flash-test", "configured": True, "modalities": ["text", "image"]},
            {"role": "vision_review", "provider": "zhipu", "model": "glm-test", "configured": False, "modalities": ["text", "image"]},
        ]
        with (
            patch("pass_bid_writing.visual_sources.ModelRouter.status", return_value=statuses),
            patch(
                "pass_bid_writing.visual_sources.ModelRouter.complete_json",
                return_value={
                    "document_summary": "河道断面图",
                    "facts": [
                        {
                            "key": "图纸名称",
                            "value": "河道标准断面图",
                            "unit": "",
                            "page": 1,
                            "figure": "HD-01",
                            "excerpt": "河道标准断面图",
                            "confidence": 0.96,
                        }
                    ],
                    "tables": [],
                    "drawings": [{"drawing_no": "HD-01", "title": "河道标准断面图", "page": 1}],
                    "conflicts": [],
                    "needs_human_review": [],
                },
            ),
        ):
            analyzed = analyze_visual_sources(state["run_id"])
        visual_fact = next(item for item in analyzed["facts"] if item["key"] == "图纸名称")
        self.assertEqual(visual_fact["source_page"], 1)
        self.assertEqual(visual_fact["source_figure"], "HD-01")
        self.assertEqual(analyzed["status"], "awaiting_confirmation")

    def test_dxf_is_rendered_to_real_preview(self) -> None:
        import ezdxf

        dxf_path = self.project / "施工总平图.dxf"
        document = ezdxf.new()
        modelspace = document.modelspace()
        modelspace.add_line((0, 0), (100, 0))
        modelspace.add_line((100, 0), (100, 50))
        modelspace.add_text("施工区域", height=5).set_placement((10, 10))
        document.saveas(dxf_path)
        preview = _render_dxf_preview(dxf_path)
        self.assertTrue(preview.exists())
        with Image.open(preview) as image:
            self.assertGreater(image.width, 500)

    def test_generated_mermaid_and_uncontrolled_basis_tail_are_removed(self) -> None:
        content = """正文说明。

```mermaid
graph TD
A --> B
```

### 质量管理组织机构图

项目经理
|
|——质量负责人——|
|                |
试验室          检验组

### 本章依据追溯

| 依据类型 | 来源 |
|---|---|
| 项目依据 | C:\\绝对路径\\招标文件.pdf |
"""
        cleaned = _sanitize_generated_structure(content)
        self.assertIn("正文说明", cleaned)
        self.assertIn("原生图形组件", cleaned)
        self.assertNotIn("graph TD", cleaned)
        self.assertNotIn("质量负责人——", cleaned)
        self.assertNotIn("绝对路径", cleaned)

    def test_bare_resource_quantities_in_tables_require_evidence(self) -> None:
        content = """| 序号 | 工种 | 计划平均人数 | 计划高峰人数 |
|---|---|---|---|
| 1 | 测量工 | 2 | 4 |
"""
        sanitized, findings = _sanitize_unsupported_claims(
            content,
            {"facts": [], "standard_clauses": []},
            {"requirements": [], "project_basis": []},
        )
        self.assertIn("| 1 | 测量工 | 【待人工确认：2】 | 【待人工确认：4】 |", sanitized)
        self.assertGreaterEqual(len(findings), 2)

    def test_verified_standard_clause_is_not_blocked(self) -> None:
        content = "施工总布置执行 SL 303-2017 第3.2.1条。"
        sanitized, findings = _sanitize_unsupported_claims(
            content,
            {
                "facts": [],
                "standard_clauses": [
                    {
                        "standard_code": "SL 303-2017",
                        "clause_id": "3.2.1",
                        "verification_status": "verified_local_source",
                    }
                ],
            },
            {"requirements": [], "project_basis": []},
        )
        self.assertEqual(content, sanitized)
        self.assertEqual([], findings)

    def test_supported_duration_is_not_blocked_and_old_placeholder_is_restored(self) -> None:
        state = {
            "facts": [{"key": "总工期", "value": "180", "unit": "日历天"}],
            "standard_clauses": [],
        }
        content = "总工期180天；原占位【待人工确认：180天的项目或规范依据】。"
        sanitized, findings = _sanitize_unsupported_claims(
            content,
            state,
            {"requirements": [], "project_basis": []},
        )
        self.assertEqual("总工期180天；原占位180天。", sanitized)
        self.assertEqual([], findings)

    def test_final_review_false_claims_are_dismissed_with_reasons(self) -> None:
        accepted, dismissed = _validate_review_issues(
            {
                "facts": [{"key": "总工期", "value": "180", "unit": "日历天"}],
                "sections": [{"content": "总工期为180日历天。"}],
                "assets": {"schedule_template_png": "schedule.png"},
                "standard_clauses": [],
            },
            [
                {
                    "title": "总工期矛盾",
                    "detail": "正文出现367天。",
                    "severity": "high",
                },
                {
                    "title": "缺少进度图",
                    "detail": "未提供横道图。",
                    "severity": "medium",
                },
            ],
        )
        self.assertEqual([], accepted)
        self.assertEqual(2, len(dismissed))


if __name__ == "__main__":
    unittest.main()
