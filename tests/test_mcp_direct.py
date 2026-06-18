from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document


class PassBidWritingMcpTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        os.environ["PASS_BID_WRITING_STORAGE_DIR"] = str(Path(self.tempdir.name) / "storage")
        self._restore_env = {}
        for key in (
            "PASS_BID_PROJECTS_DIR",
            "PASS_BID_VISION_API_KEY",
            "DASHSCOPE_API_KEY",
            "MOONSHOT_API_KEY",
            "ARK_API_KEY",
            "VOLCENGINE_API_KEY",
            "PASS_BID_VISION_BASE_URL",
        ):
            self._restore_env[key] = os.environ.get(key)
            os.environ.pop(key, None)
        self.addCleanup(self._restore_environment)

    def _restore_environment(self) -> None:
        for key, value in self._restore_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def test_case_ingest_matrix_docx_and_check(self) -> None:
        from pass_bid_writing import mcp_server

        root = Path(self.tempdir.name)
        tender = root / "tender.md"
        bid = root / "passed.md"
        tender.write_text(
            "\n".join(
                [
                    "本项目技术标采用通过制，投标文件应实质性响应招标文件要求。",
                    "施工组织设计应包含工期、质量、安全文明施工、环境保护、资源配置。",
                    "项目涉及河道治理、围堰导流、度汛和排水措施。",
                ]
            ),
            encoding="utf-8",
        )
        bid.write_text(
            "\n".join(
                [
                    "一、编制说明及工程概况",
                    "本技术标完全响应招标文件要求。",
                    "二、施工进度计划及保证措施",
                    "合理安排工期节点。",
                    "三、河道治理、围堰导流及度汛措施",
                    "设置围堰导流和排水体系，确保安全度汛。",
                ]
            ),
            encoding="utf-8",
        )

        ingested = mcp_server.writing_ingest_case_pair(
            tender_file=str(tender),
            bid_file=str(bid),
            title="河道治理通过制案例",
            project_type="河道治理",
            tags="water,pass",
        )
        self.assertGreaterEqual(ingested["summary"]["requirement_count"], 2)

        requirements = mcp_server.writing_extract_tender_requirements(tender_file=str(tender))
        matrix = mcp_server.writing_build_response_matrix(requirements=requirements)["matrix"]
        self.assertGreaterEqual(len(matrix), 2)

        outline = mcp_server.writing_generate_outline(
            requirements=requirements,
            project_type="河道治理",
        )["outline"]
        docx = mcp_server.writing_generate_docx(
            title="测试项目通过制技术标",
            outline=outline,
            requirements=requirements,
            response_matrix=matrix,
            output_name="test-draft",
            tender_path=str(tender),
            visual_qa=False,
        )
        self.assertTrue(Path(docx["docx_path"]).exists())
        self.assertGreaterEqual(docx["table_count"], 1)
        self.assertEqual(docx["content_status"], "placeholder_needs_section_text")

        compliance = mcp_server.writing_check_draft_compliance(
            draft_docx=docx["docx_path"],
            requirements=requirements,
        )
        self.assertIn("status", compliance)

    def test_generate_docx_accepts_wrapped_tool_outputs_and_compliance_ignores_matrix_table(self) -> None:
        from pass_bid_writing import mcp_server

        requirement_text = "必须设置龙门吊专项安全防护和沉降观测记录"
        requirements = {
            "detected_mode": "pass_fail",
            "requirement_count": 1,
            "requirements": [
                {
                    "id": "REQ-001",
                    "category": "safety",
                    "requirement": requirement_text,
                    "source_line": 1,
                    "suggested_section": "安全生产、文明施工及应急措施",
                    "status": "needs_response",
                }
            ],
        }
        outline_payload = mcp_server.writing_generate_outline(requirements=requirements, title="链路兼容测试")
        matrix_payload = mcp_server.writing_build_response_matrix(
            requirements=requirements,
            outline=outline_payload,
        )

        docx = mcp_server.writing_generate_docx(
            title="链路兼容测试",
            outline=outline_payload,
            requirements=requirements,
            response_matrix=matrix_payload,
            output_name="wrapped-payload-draft",
            visual_qa=False,
        )
        self.assertEqual(docx["content_status"], "placeholder_needs_section_text")

        doc = Document(docx["docx_path"])
        table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn(requirement_text, table_text)

        compliance = mcp_server.writing_check_draft_compliance(
            draft_docx=docx["docx_path"],
            requirements=requirements,
        )
        self.assertEqual(compliance["checked_text_scope"], "paragraphs_only")
        self.assertEqual(compliance["missing_count"], 1)
        self.assertEqual(compliance["covered_count"], 0)

    def test_generate_docx_renders_cover_toc_header_footer_and_markdown_table(self) -> None:
        from pass_bid_writing import mcp_server

        result = mcp_server.writing_generate_docx(
            title="新建泵站工程通过制技术标",
            sections=[
                {
                    "level": 1,
                    "title": "第一章 编制说明",
                    "content": "\n".join(
                        [
                            "本章说明编制依据和响应原则。",
                            "| 序号 | 响应内容 | 状态 |",
                            "| --- | --- | --- |",
                            "| 1 | 工期要求 | 已响应 |",
                            "| 2 | 质量要求 | 已响应 |",
                            "## 1.1 编制依据",
                            "严格响应招标文件要求。",
                        ]
                    ),
                }
            ],
            output_name="markdown-table-draft",
            visual_qa=False,
            layout_profile_id="",
        )
        self.assertEqual(result["content_status"], "section_text_provided")

        doc = Document(result["docx_path"])
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertIn("新建泵站工程通过制技术标", text)
        self.assertIn("目录", text)
        self.assertGreaterEqual(len(doc.tables), 1)
        table_text = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(table_text[0], ["序号", "响应内容", "状态"])
        self.assertEqual(table_text[1], ["1", "工期要求", "已响应"])

        header_text = doc.sections[0].header.paragraphs[0].text
        footer_text = doc.sections[0].footer.paragraphs[0].text
        self.assertIn("新建泵站工程通过制技术标", header_text)
        self.assertIn("第 ", footer_text)

    def test_generate_docx_applies_saved_layout_profile(self) -> None:
        from pass_bid_writing import db, mcp_server
        from pass_bid_writing.config import ensure_storage_dirs, get_settings

        settings = get_settings()
        ensure_storage_dirs(settings)
        db.init_db(settings.database_path)
        with db.db_session(settings.database_path) as conn:
            layout_profile_id = db.create_layout_profile(
                conn,
                source_path="layout-source.pdf",
                source_kind="accepted_bid",
                provider="test",
                model="test-model",
                status="ready",
                profile={
                    "cover": {"subtitle": "已通过样式技术标", "bottom_text": "测试投标单位"},
                    "headers_footers": {"header": "样式化页眉"},
                },
                render={},
            )

        result = mcp_server.writing_generate_docx(
            title="样式应用测试",
            sections={"第一章 测试": "正文内容。"},
            output_name="layout-profile-draft",
            visual_qa=False,
            layout_profile_id=str(layout_profile_id),
        )
        doc = Document(result["docx_path"])
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertIn("已通过样式技术标", text)
        self.assertIn("测试投标单位", text)
        self.assertIn("样式化页眉", doc.sections[0].header.paragraphs[0].text)
        self.assertEqual(result["style_config"]["cover_color"], "1F4E79")

    def test_visual_profile_style_values_drive_table_and_cover_formatting(self) -> None:
        from pass_bid_writing import db, mcp_server
        from pass_bid_writing.config import ensure_storage_dirs, get_settings

        settings = get_settings()
        ensure_storage_dirs(settings)
        db.init_db(settings.database_path)
        with db.db_session(settings.database_path) as conn:
            layout_profile_id = db.create_layout_profile(
                conn,
                source_path="styled-layout.pdf",
                source_kind="accepted_bid",
                provider="test",
                model="test-model",
                status="ready",
                profile={
                    "cover": {
                        "subtitle": "蓝绿波浪封面",
                        "bottom_text": "台州水利投标文件",
                        "style": "标题为深蓝色，封面有蓝绿色波浪装饰",
                    },
                    "fonts": {"body": "正文仿宋", "heading": "标题黑体蓝色"},
                    "tables": {"style": "蓝底表头，表格正文宋体"},
                    "style_rules": ["表头底色为#00A65A", "标题使用深蓝色"],
                },
                render={},
            )

        result = mcp_server.writing_generate_docx(
            title="视觉样式应用测试",
            sections=[
                {
                    "title": "第一章 表格测试",
                    "content": "\n".join(
                        [
                            "| 名称 | 要求 |",
                            "| --- | --- |",
                            "| 工期 | 满足招标文件 |",
                        ]
                    ),
                }
            ],
            output_name="visual-style-draft",
            visual_qa=False,
            layout_profile_id=layout_profile_id,
        )
        self.assertEqual(result["style_config"]["table_header_fill"], "00A65A")
        self.assertEqual(result["style_config"]["heading_color"], "1F4E79")
        self.assertEqual(result["style_config"]["body_font"][1], "仿宋")
        self.assertIn("蓝绿波浪封面", "\n".join(p.text for p in Document(result["docx_path"]).paragraphs))

    def test_visual_profile_maps_absent_header_heading_color_and_numbering(self) -> None:
        from pass_bid_writing import db, mcp_server
        from pass_bid_writing.config import ensure_storage_dirs, get_settings

        settings = get_settings()
        ensure_storage_dirs(settings)
        db.init_db(settings.database_path)
        with db.db_session(settings.database_path) as conn:
            layout_profile_id = db.create_layout_profile(
                conn,
                source_path="xiaozhi-layout.pdf",
                source_kind="accepted_bid",
                provider="test",
                model="test-model",
                status="ready",
                profile={
                    "cover": {"style": "封面保留蓝绿波浪装饰"},
                    "fonts": {"body_font": "宋体", "heading_font": "黑体，加粗，黑色"},
                    "heading_hierarchy": ["第一章/第一节/一、/1、"],
                    "numbering": {"levels": ["第一章", "第一节", "一、", "1、"]},
                    "headers_footers": {"header": "无"},
                    "tables": {"header_fill": "#4472C4"},
                    "style_rules": ["封面蓝绿波浪装饰，标题黑体、加粗、黑色"],
                },
                render={},
            )

        result = mcp_server.writing_generate_docx(
            title="小芝镇样式映射测试",
            sections=[
                {
                    "level": 1,
                    "title": "1 编制说明",
                    "content": "\n".join(
                        [
                            "## 1.1 编制依据",
                            "1. 招标文件。",
                            "2. 施工图纸。",
                            "",
                            "## 1.2 施工部署",
                            "1. 围堰导流。",
                        ]
                    ),
                }
            ],
            output_name="xiaozhi-style-draft",
            visual_qa=False,
            layout_profile_id=layout_profile_id,
        )

        doc = Document(result["docx_path"])
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertEqual(result["style_config"]["heading_color"], "000000")
        self.assertEqual(result["style_config"]["table_header_fill"], "4472C4")
        self.assertEqual(result["style_config"]["header_text"], "")
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "")
        self.assertIn("第一章 编制说明", text)
        self.assertIn("第一节 编制依据", text)
        self.assertIn("第二节 施工部署", text)
        self.assertIn("1、招标文件。", text)
        self.assertIn("2、施工图纸。", text)
        self.assertIn("1、围堰导流。", text)

    def test_project_folder_aliases_and_new_tender_prepare(self) -> None:
        from pass_bid_writing import mcp_server

        root = Path(self.tempdir.name)
        projects = root / "projects"
        passed = projects / "passed" / "河道治理已通过"
        unpassed = projects / "unpassed" / "新河道项目"
        passed.mkdir(parents=True)
        unpassed.mkdir(parents=True)
        (passed / "招标文件.md").write_text(
            "本项目技术标采用通过制，应包含工期、质量、安全文明施工和环境保护。",
            encoding="utf-8",
        )
        (passed / "已通过技术标.md").write_text(
            "\n".join(
                [
                    "一、编制说明及工程概况",
                    "本技术标完全响应招标文件。",
                    "二、围堰导流及度汛措施",
                    "设置围堰导流和排水体系。",
                ]
            ),
            encoding="utf-8",
        )
        (unpassed / "招标文件.md").write_text(
            "通过制技术标应响应工期、质量、安全、环保、河道治理和度汛要求。",
            encoding="utf-8",
        )

        scanned = mcp_server.writing_scan_projects(root=str(projects))
        self.assertEqual(scanned["passed_case_count"], 1)
        self.assertEqual(scanned["new_tender_count"], 1)
        self.assertTrue(scanned["passed_cases"][0]["ready"])
        self.assertTrue(scanned["new_tenders"][0]["ready"])

        ingested = mcp_server.writing_ingest_passed_case(
            project_dir=str(passed),
            auto_visual=False,
        )
        self.assertGreaterEqual(ingested["summary"]["outline_count"], 2)
        self.assertIsNone(ingested["layout_profile"])

        prepared = mcp_server.writing_prepare_new_tender(
            project_dir=str(unpassed),
            auto_visual=False,
        )
        self.assertGreaterEqual(prepared["requirements"]["requirement_count"], 1)
        self.assertGreaterEqual(len(prepared["response_matrix"]), 1)
        self.assertGreaterEqual(len(prepared["similar_cases"]), 1)
        self.assertGreaterEqual(len(prepared["recommended_search_queries"]), 1)
        self.assertIsNone(prepared["new_tender_layout_profile_id"])
        self.assertTrue(prepared["outputs_dir"].endswith("outputs"))

    def test_layout_profile_without_api_key_renders_and_degrades(self) -> None:
        from pass_bid_writing import mcp_server

        pdf = Path(self.tempdir.name) / "layout-sample.pdf"
        self._write_sample_pdf(pdf)

        result = mcp_server.writing_extract_layout_profile(
            file_path=str(pdf),
            source_kind="accepted_bid",
        )
        self.assertEqual(result["status"], "not_configured")
        self.assertEqual(result["render"]["status"], "ready")
        self.assertTrue(result["layout_profile_id"])
        self.assertTrue(Path(result["render"]["pages"][0]["image_path"]).exists())

    def test_layout_profile_parses_mock_vision_json(self) -> None:
        from pass_bid_writing import mcp_server

        os.environ["PASS_BID_VISION_API_KEY"] = "test-key"
        os.environ["PASS_BID_VISION_BASE_URL"] = "https://example.test/v1"
        pdf = Path(self.tempdir.name) / "mock-vision.pdf"
        self._write_sample_pdf(pdf)

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def read(self):
                return json_bytes(
                    '{"choices":[{"message":{"content":"{\\"cover\\":{\\"subtitle\\":\\"技术标\\"},'
                    '\\"heading_hierarchy\\":[\\"一级标题居中\\"],\\"visual_findings\\":[]}"}}]}'
                )

        with patch("urllib.request.urlopen", return_value=FakeResponse()):
            result = mcp_server.writing_extract_layout_profile(
                file_path=str(pdf),
                source_kind="accepted_bid",
            )
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["profile"]["cover"]["subtitle"], "技术标")
        self.assertTrue(result["layout_profile_id"])

    def _write_sample_pdf(self, path: Path) -> None:
        import fitz  # type: ignore

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Pass bid technical document layout sample")
        page.insert_text((72, 110), "Cover / TOC / heading / footer visual test")
        doc.save(str(path))
        doc.close()


def json_bytes(value: str) -> bytes:
    return value.encode("utf-8")


if __name__ == "__main__":
    unittest.main()
